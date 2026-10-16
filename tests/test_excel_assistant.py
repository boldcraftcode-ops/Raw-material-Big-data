import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.db import Database
from backend.excel_assistant import column_advice, sample_advice
from backend.parsers import parse_file

A={'X-Workspace-Token':'E'*64}
B={'X-Workspace-Token':'F'*64}
HEAD=['Tanggal','Pemasok','Material','Harga','Mata Uang','Unit','Basis']
DATA=['2026-10-15','Pemasok Uji','Bokar',15000,'IDR','kg','basah']


@pytest.fixture
def client(app_factory):
    with TestClient(app_factory(seed_demo=False,start_worker=False,seed_sources=False)) as client:
        yield client


def workbook(rows,second=None):
    book=Workbook();sheet=book.active;sheet.title='Pembelian'
    for row in rows:sheet.append(row)
    if second:
        sheet=book.create_sheet('Catatan')
        for row in second:sheet.append(row)
    out=io.BytesIO();book.save(out);return out.getvalue()


def drain(client):
    for _ in range(20):
        if not client.app.state.service.process_next():return
    raise AssertionError('Jobs not drained')


def upload(client,rows=None,data=None,name='contoh.xlsx',headers=A):
    data=data if data is not None else workbook(rows or [HEAD,DATA])
    result=client.post('/api/uploads',headers=headers,files={'file':(name,data)},data={'consent':'true'})
    assert result.status_code==202,result.text
    version=result.json()['version']['id'];drain(client)
    preview=client.get(f'/api/versions/{version}/preview',headers=headers)
    assert preview.status_code==200,preview.text
    return version,preview.json(),data


def config(preview,**extra):
    return {'sheet':preview['sheet'],'fields':preview['suggested_mapping'],
            'options':{'number_format':'id-ID','date_format':'iso','drc_scale':'percent','default_price_kind':'dilaporkan'},
            'header_row':preview['assistant']['header_row'],'header_depth':preview['assistant']['header_depth'],
            'scope_confirmed':True,**extra}


def publish(client,version,mapping,headers=A):
    result=client.post(f'/api/versions/{version}/mapping',headers=headers,json=mapping)
    assert result.status_code==202,result.text
    drain(client)
    return result.json()['revision']


def remember(client,version,revision,headers=A):
    return client.post(f'/api/versions/{version}/mapping/remember',headers=headers,json={'revision':revision})


def test_leading_titles_and_blank_rows_keep_real_coordinates(client):
    version,preview,raw=upload(client,[['Laporan pembelian karet'],[],['Periode Oktober'],HEAD,DATA])
    assert preview['assistant']['mode']=='local_rules' and preview['assistant']['external_ai_used'] is False
    assert preview['assistant']['header_row']==4 and preview['rows'][0]['row_no']==5
    assert preview['row_count']==1 and preview['assistant']['requires_confirmation']
    publish(client,version,config(preview))
    observation=client.get('/api/observations?dataset=live',headers=A).json()['items'][0]
    evidence=client.get('/api/evidence/'+observation['id'],headers=A).json()
    assert evidence['row_no']==5 and evidence['header_row']==4
    assert client.get(f'/api/versions/{version}/raw',headers=A).content==raw


def test_two_level_header_combines_groups_without_formula_execution(client):
    rows=[['Laporan karet'],[],['Tanggal','Pemasok','Bahan','Harga',None,None,'Basis'],
          [None,None,None,'Nilai','Mata Uang','Unit',None],DATA]
    version,preview,_=upload(client,rows)
    assert preview['assistant']['header_row']==3 and preview['assistant']['header_depth']==2
    assert preview['columns'][3]['label']=='Harga / Nilai'
    assert preview['suggested_mapping']['price']=='c3' and preview['suggested_mapping']['currency']=='c4'
    publish(client,version,config(preview))
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==1


def test_recommend_sheet_with_known_purchase_structure(client):
    book=Workbook();book.active.title='Catatan';book.active.append(['Keterangan']);book.active.append(['Tidak ada transaksi di sini'])
    sheet=book.create_sheet('Transaksi');sheet.append(HEAD);sheet.append(DATA)
    out=io.BytesIO();book.save(out)
    _,preview,_=upload(client,data=out.getvalue())
    assert preview['sheet']=='Transaksi' and len(preview['sheets'])==2
    assert any('satu lembar' in message for message in preview['warnings'])


def test_csv_preamble_and_quoted_multiline_cells_preserve_line_pointers(client):
    raw=b'"Laporan\nKaret"\nTanggal,Pemasok,Material,Harga,Mata Uang,Unit,Basis\n2026-10-15,"Pemasok\nBaru",Bokar,15000,IDR,kg,basah\n'
    version,preview,_=upload(client,data=raw,name='judul.csv')
    assert preview['assistant']['header_row']==3
    assert preview['rows'][0]['row_no']==4 and preview['rows'][0]['line_end']==5
    publish(client,version,config(preview))
    assert client.get(f'/api/versions/{version}/raw',headers=A).content==raw


def test_semicolon_with_title_is_detected(client):
    raw=('Laporan pembelian\n'+ ';'.join(HEAD)+'\n'+';'.join(map(str,DATA))+'\n').encode()
    _,preview,_=upload(client,data=raw,name='delimiter.csv')
    assert preview['assistant']['header_row']==2 and len(preview['columns'])==7


def test_stock_not_silently_classified_or_published_as_purchases(client):
    _,preview,_=upload(client,[['Tanggal','Material','Saldo kg'],['2026-10-15','Bokar',200]])
    assert preview['assistant']['classification']=='other_or_uncertain'
    assert 'price' in preview['assistant']['missing_required'] and 'supplier' in preview['assistant']['missing_required']
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==0


def test_duplicate_header_matches_are_not_silently_picked():
    columns=[{'key':'c0','label':'Harga'},{'key':'c1','label':'Harga Beli'}]
    advice=column_advice(columns)
    assert 'price' not in advice['suggested_mapping'] and advice['conflicts'][0]['field']=='price'


def test_ambiguous_dates_and_drc_scale_are_not_guessed():
    columns=[{'key':'c0','label':'Tanggal'},{'key':'c1','label':'KKK'}]
    advice=sample_advice(columns,[{'cells':{'c0':'02/03/2026','c1':.6}}])
    assert 'date_format' not in advice['option_hints'] and 'drc_scale' not in advice['option_hints']
    assert any('ambigu' in note for note in advice['notes'])


def test_unambiguous_format_hints_are_explained_not_forced():
    advice=sample_advice([{'key':'c0','label':'Tanggal'},{'key':'c1','label':'Harga'}],[{'cells':{'c0':'15/10/2026','c1':'15.000,50'}}])
    assert advice['option_hints']=={'date_format':'dmy','number_format':'id-ID'}
    assert 'default_currency' not in advice['option_hints'] and 'default_basis' not in advice['option_hints']


def test_sample_check_is_bounded_readonly_and_does_not_publish(client):
    version,preview,_=upload(client,[HEAD]+[DATA]*60)
    mapping=config(preview,scope_confirmed=False)
    response=client.post(f'/api/versions/{version}/mapping/check',headers=A,json=mapping)
    assert response.status_code==200,response.text
    report=response.json()
    assert report['sampled_rows']==50 and report['total_rows']==60 and report['valid_rows']==50
    assert report['published'] is False and report['full_validation'] is False
    with client.app.state.db.connect() as conn:
        assert conn.execute('SELECT COUNT(*) FROM mappings').fetchone()[0]==0
        assert conn.execute('SELECT COUNT(*) FROM observations').fetchone()[0]==0
        assert conn.execute('SELECT COUNT(*) FROM excel_templates').fetchone()[0]==0
    assert client.get(f'/api/versions/{version}',headers=A).json()['status']=='awaiting_mapping'


def test_invalid_sample_lists_original_rows_without_publication(client):
    bad=DATA.copy();bad[3]='bukan angka'
    version,preview,_=upload(client,[['Judul'],HEAD,bad])
    result=client.post(f'/api/versions/{version}/mapping/check',headers=A,json=config(preview)).json()
    assert result['invalid_rows']==1 and result['issues'][0]['row_no']==3
    assert any(error['field']=='price' for error in result['issues'][0]['errors'])


def test_explicit_header_override_is_readonly_and_old_evidence_survives(client):
    alternate=['Tgl','Nama Supplier','Jenis Karet','Harga Beli','Valuta','Satuan Harga','Weight Basis']
    version,preview,raw=upload(client,[HEAD,DATA,alternate,DATA])
    publish(client,version,config(preview))
    old=client.get('/api/observations?dataset=live',headers=A).json()['items']
    assert len(old)==2
    override=client.get(f'/api/versions/{version}/preview?sheet=Pembelian&header_row=3',headers=A).json()
    assert override['row_count']==1 and override['rows'][0]['row_no']==4
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==2
    publish(client,version,config(override))
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==1
    evidence=client.get('/api/evidence/'+old[0]['id'],headers=A).json()
    assert evidence['raw_values'][0]['label']=='Tanggal' and evidence['header_row']==1
    assert not evidence['is_active']
    assert client.get(f'/api/versions/{version}/raw',headers=A).content==raw


def test_memory_requires_explicit_approval_and_no_cell_values_are_saved(client):
    custom=HEAD.copy();custom[1]='Mitra Pasok'
    version,preview,_=upload(client,[custom,DATA])
    assert not preview['suggested_mapping'].get('supplier')
    mapping=config(preview);mapping['fields']['supplier']='c1'
    assert remember(client,version,1).status_code==400
    revision=publish(client,version,mapping)
    with client.app.state.db.connect() as conn:assert conn.execute('SELECT COUNT(*) FROM excel_templates').fetchone()[0]==0
    result=remember(client,version,revision)
    assert result.status_code==200,result.text
    assert remember(client,version,revision).json()['id']==result.json()['id']
    with client.app.state.db.connect() as conn:
        rows=conn.execute('SELECT * FROM excel_templates').fetchall();assert len(rows)==1
        assert 'Pemasok Uji' not in json.dumps(dict(rows[0])) and '15000' not in json.dumps(dict(rows[0]))


def learned(client):
    custom=HEAD.copy();custom[1]='Mitra Pasok'
    version,preview,raw=upload(client,[custom,DATA])
    mapping=config(preview);mapping['fields']['supplier']='c1'
    revision=publish(client,version,mapping)
    template=remember(client,version,revision).json()['id']
    return version,template,raw


def test_approved_memory_follows_reordered_columns_not_filename(client):
    learned(client)
    header=['Harga','Basis','Mitra Pasok','Tanggal','Material','Unit','Mata Uang']
    row=[16000,'basah','Pemasok Lain','2026-10-16','Bokar','kg','IDR']
    _,preview,_=upload(client,[header,row],name='format-baru.xlsx')
    profile=preview['assistant']['recalled_profile']
    assert profile and profile['fields']['supplier']=='c2' and profile['fields']['date']=='c3'
    assert preview['mapping'] is None
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==1


def test_owner_scope_applies_to_preview_check_memory_and_forget(client):
    version,template,raw=learned(client)
    _,preview,_=upload(client,data=raw,name='milik-b.xlsx',headers=B)
    assert preview['assistant']['recalled_profile'] is None
    assert client.get(f'/api/versions/{version}/preview?header_row=1',headers=B).status_code==404
    assert client.post(f'/api/versions/{version}/mapping/check',headers=B,json=config(preview)).status_code==404
    assert remember(client,version,1,headers=B).status_code==404
    assert client.delete('/api/excel-templates/'+template,headers=B).status_code==404


def test_forget_does_not_change_raw_or_observations(client):
    version,template,raw=learned(client)
    assert client.delete('/api/excel-templates/'+template,headers=A).status_code==200
    assert client.get(f'/api/versions/{version}/preview',headers=A).json()['assistant']['recalled_profile'] is None
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==1
    assert client.get(f'/api/versions/{version}/raw',headers=A).content==raw


def test_deleting_source_removes_learned_profile(client):
    version,_,_=learned(client)
    document=client.get(f'/api/versions/{version}',headers=A).json()['document_id']
    assert client.delete('/api/documents/'+document,headers=A).status_code==200
    with client.app.state.db.connect() as conn:assert conn.execute('SELECT COUNT(*) FROM excel_templates').fetchone()[0]==0


def test_zero_valid_rows_cannot_teach_a_template(client):
    invalid=DATA.copy();invalid[3]='invalid'
    version,preview,_=upload(client,[HEAD,invalid])
    revision=publish(client,version,config(preview))
    assert remember(client,version,revision).status_code==400


def test_duplicate_labels_prevent_unsafe_memory(client):
    header=HEAD+['Harga']
    version,preview,_=upload(client,[header,DATA+[17000]])
    mapping=config(preview);mapping['fields']['price']='c3'
    revision=publish(client,version,mapping)
    assert remember(client,version,revision).status_code==400


def test_invalid_header_parameters_rejected(client):
    version,_,_=upload(client)
    assert client.get(f'/api/versions/{version}/preview?header_row=0',headers=A).status_code==422
    assert client.get(f'/api/versions/{version}/preview?header_depth=3',headers=A).status_code==422
    assert client.get(f'/api/versions/{version}/preview?header_row=49',headers=A).status_code==400


def test_memory_survives_database_reopen(client):
    _,template,_=learned(client)
    database=client.app.state.db
    if database.kind!='sqlite':return
    reopened=Database(database.root)
    with reopened.connect() as conn:assert conn.execute('SELECT id FROM excel_templates').fetchone()['id']==template


def test_unsupported_formula_result_never_becomes_invented_price(client):
    row=DATA.copy();row[3]='=SUM(10000,5000)'
    version,preview,_=upload(client,[HEAD,row])
    assert preview['rows'][0]['cells']['c3'] is None
    result=client.post(f'/api/versions/{version}/mapping/check',headers=A,json=config(preview)).json()
    assert result['invalid_rows']==1 and result['valid_rows']==0
    assert any('Formula' in warning for warning in preview['warnings'])
