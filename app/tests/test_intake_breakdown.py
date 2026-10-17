"""Synthetic MVP cases. These are rule/structure tests, not a real-workbook accuracy benchmark."""
import copy
import io
import json
from datetime import datetime
from zipfile import ZipFile
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from backend.intake_breakdown import build_breakdown

A={'X-Workspace-Token':'intake_mvp_owner_a_'+'a'*48}
B={'X-Workspace-Token':'intake_mvp_owner_b_'+'b'*48}
STOCK=b'Tanggal,Material,Stok,Gudang\n2026-10-15,Lateks,100,G1\n2026-10-16,Bokar,0,G2\n'


def workbook(sheets):
    book=Workbook();book.remove(book.active)
    for name,rows in sheets.items():
        sheet=book.create_sheet(name)
        for row in rows:sheet.append(row)
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def tables(report):return [table for sheet in report['sheets'] for table in sheet['tables']]


@pytest.mark.parametrize('header,row,category',[
 ('Tanggal,Pemasok,Material,Harga','2026-10-15,Koperasi,Lateks,14000','purchase'),
 ('Tanggal,Pelanggan,Material,Harga','2026-10-15,Pabrik,Bokar,14000','sales'),
 ('Tanggal,Material,Harga','2026-10-15,TSR20,1.8','price'),
 ('Tanggal,Material,Stok,Gudang','2026-10-15,Lateks,0,G1','inventory'),
 ('Tanggal,Material,Barang masuk,Barang keluar','2026-10-15,Lateks,100,2','inventory'),
 ('Tanggal,Material,Hasil produksi','2026-10-15,Bokar,120','production'),
 ('Tanggal,Material,Kadar abu','2026-10-15,Lateks,0.5','quality'),
 ('Tanggal,ID sampel,Hasil uji','2026-10-15,S1,60','quality'),
 ('Tanggal,No surat jalan,Tujuan','2026-10-15,SJ1,Medan','logistics'),
 ('Tanggal,Curah hujan,Suhu','2026-10-15,10,28','weather'),
 ('Lot,Asal,Material','LOT1,Kebun A,Lateks','traceability'),
])
def test_business_categories_are_combinations_of_column_roles(header,row,category):
    report=build_breakdown((header+'\n'+row+'\n').encode(),'rubber.csv')
    assert report['status']=='ready' and report['classification']==category
    assert tables(report)[0]['classification']['matches'][0]['category']==category
    assert report['external_ai_used'] is False and report['published_as_transactions'] is False


def test_unknown_does_not_invent_purchase_fields_or_identity():
    report=build_breakdown(b'Alpha,Beta\na,b\nc,d\n','unknown.csv')
    table=tables(report)[0]
    assert report['classification']=='unknown' and table['header_detection']=='structural_guess'
    assert all(column['role']=='unknown' for column in table['columns'])
    assert table['record_count']==2


def test_mixed_business_signals_are_not_forced_to_purchase():
    report=build_breakdown(b'Tanggal,Material,Pemasok,Harga,Stok\n2026-10-15,Lateks,Koperasi,14000,500\n','mixed.csv')
    assert report['classification']=='mixed'
    assert {item['category'] for item in tables(report)[0]['classification']['matches']}=={'purchase','inventory'}


def test_workbook_breaks_down_per_sheet_and_keeps_physical_coordinates():
    raw=workbook({'Stok':[['Laporan sintetis'],[],['Tanggal','Material','Stok','Gudang'],['2026-10-15','Lateks',0,'G1']],
                  'Produksi':[['Tanggal','Material','Hasil produksi'],['2026-10-15','Bokar',120]]})
    report=build_breakdown(raw,'mixed.xlsx');stock,production=tables(report)
    assert report['classification']=='mixed' and report['summary']['tables_detected']==2
    assert stock['header_row']==3 and stock['preview'][0]['row_no']==4
    assert stock['classification']['category']=='inventory' and production['classification']['category']=='production'
    assert stock['columns'][2]['source_ref']=='Stok!C3'
    assert report['summary']['unassigned_nonempty_rows']==1


def test_separate_tables_and_repeated_headers_are_not_data_rows():
    raw=b'Catatan laporan\nTanggal,Material,Stok\n2026-10-15,Lateks,1\n\nTanggal,Material,Hasil produksi\n2026-10-16,Bokar,2\n'
    report=build_breakdown(raw,'sections.csv')
    assert [table['header_row'] for table in tables(report)]==[2,5]
    assert [table['record_count'] for table in tables(report)]==[1,1]
    repeat=build_breakdown(b'Tanggal,Material,Stok\n2026-10-15,Lateks,1\nTanggal,Material,Stok\n2026-10-16,Bokar,2\n','repeat.csv')
    assert [table['header_row'] for table in tables(repeat)]==[1,3]
    assert sum(table['record_count'] for table in tables(repeat))==2


def test_side_by_side_tables_require_a_fully_blank_separator():
    raw=workbook({'Data':[['Tanggal','Material','Stok',None,'Tanggal','Material','Hasil produksi'],
                          ['2026-10-15','Lateks',0,None,'2026-10-15','Bokar',120]]})
    report=build_breakdown(raw,'side.xlsx');one,two=tables(report)
    assert one['start_column']==1 and one['end_column']==3
    assert two['start_column']==5 and two['columns'][0]['id']=='c4'
    assert two['columns'][0]['source_ref']=='Data!E1'
    assert report['summary']['data_records_scanned']==2 and report['summary']['physical_rows_scanned']==2


def test_two_level_headers_are_source_bound():
    raw=workbook({'Data':[['Tanggal','Pemasok','Material','Harga',None,None],
                          [None,None,None,'Nilai','Mata Uang','Unit'],
                          ['2026-10-15','Koperasi','Lateks',14000,'IDR','kg']]})
    table=tables(build_breakdown(raw,'two.xlsx'))[0]
    assert table['header_depth']==2 and table['header_end_row']==2
    assert table['columns'][3]['role']=='price' and table['preview'][0]['row_no']==3
    assert table['record_count']==1


def test_csv_bom_semicolon_and_multiline_physical_pointers():
    raw=('\ufeffTanggal;Material;Stok;Catatan\r\n2026-10-15;Lateks;0;"baris satu\r\nbaris dua"\r\n2026-10-16;Bokar;1;oke\r\n').encode()
    report=build_breakdown(raw,'semicolon.csv');table=tables(report)[0]
    assert table['preview'][0]['row_no']==2 and table['preview'][0]['line_end']==3
    assert table['preview'][1]['row_no']==4
    assert report['summary']['physical_rows_scanned']==4
    assert table['record_count']==2


def test_native_numbers_text_numbers_leading_zero_blanks_and_dates_are_distinct():
    raw=workbook({'Types':[['Tanggal','Material','Stok','ID','Ambigu','KKK','Campuran','Boolean'],
                           [datetime(2026,10,15),'Lateks',0,'00123','01/02/2026','60%',1,False],
                           [datetime(2026,10,16),'Bokar',None,'00005','02/03/2026','0,6%','teks',True]]})
    columns=tables(build_breakdown(raw,'types.xlsx'))[0]['columns']
    assert columns[0]['observed_type']=='date'
    assert columns[2]['observed_type']=='number' and columns[2]['empty_count']==1
    assert columns[2]['samples'][0]['value']==0
    assert columns[3]['observed_type']=='code_text' and columns[3]['samples'][0]['value']=='00123'
    assert columns[4]['observed_type']=='date_text_ambiguous'
    assert columns[5]['observed_type']=='percent_text'
    assert columns[6]['observed_type']=='mixed'
    assert columns[7]['observed_type']=='boolean'
    assert columns[4]['role']=='unknown'  # Do not infer business meaning solely from a value pattern.


def test_all_scanned_rows_contribute_to_type_counts_not_only_preview():
    raw=b'Tanggal,Material,Stok\n'+b'2026-10-15,Lateks,1\n'*30+b'2026-10-16,Lateks,belum tersedia\n'
    table=tables(build_breakdown(raw,'types.csv'))[0]
    assert len(table['preview'])==4 and table['record_count']==31
    assert table['columns'][2]['type_counts']=={'number_text':30,'text':1}
    assert table['columns'][2]['observed_type']=='mixed'


def test_row_budget_marks_partial_without_extrapolating(monkeypatch):
    monkeypatch.setattr('backend.intake_breakdown.MAX_ROWS',4)
    raw=b'Tanggal,Material,Stok\n'+b'2026-10-15,Lateks,1\n'*9
    report=build_breakdown(raw,'bound.csv')
    assert report['status']=='partial' and not report['sheets'][0]['scan_complete']
    assert tables(report)[0]['record_count']==3
    assert report['summary']['physical_rows_scanned']==4


def test_column_budget_and_unassigned_ragged_cells_are_reported(monkeypatch):
    monkeypatch.setattr('backend.intake_breakdown.MAX_COLUMNS',3)
    report=build_breakdown(STOCK,'wide.csv')
    assert report['status']=='partial' and len(tables(report)[0]['columns'])==3
    ragged=build_breakdown(b'Tanggal,Stok\n2026-10-15,1\n2026-10-16,2\n2026-10-17,3\n2026-10-18,4,tambahan\n','ragged.csv')
    assert ragged['summary']['unassigned_nonempty_cells']>=1


def test_sheet_budget_and_hidden_sheets_are_explicit(monkeypatch):
    monkeypatch.setattr('backend.intake_breakdown.MAX_SHEETS',1)
    raw=workbook({'Satu':[['Tanggal','Stok'],['2026-10-15',1]],'Dua':[['Tanggal','Stok'],['2026-10-15',2]]})
    report=build_breakdown(raw,'many.xlsx')
    assert report['status']=='partial' and report['summary']['sheets_total']==2 and report['summary']['sheets_scanned']==1
    book=Workbook();book.active.title='Visible';book.active.append(['Tanggal','Stok']);book.active.append(['2026-10-15',1])
    book.create_sheet('Hidden').sheet_state='hidden';out=io.BytesIO();book.save(out)
    hidden=build_breakdown(out.getvalue(),'hidden.xlsx')
    assert hidden['status']=='partial' and 'tersembunyi' in ' '.join(hidden['notes'])


def test_table_budget_and_unassigned_rows_are_explicit():
    raw=b'\n'.join([b'Tanggal,Stok\n2026-10-15,1\n']*7)
    report=build_breakdown(raw,'tables.csv')
    assert len(tables(report))==5 and report['status']=='partial'
    assert report['summary']['unassigned_nonempty_rows']>0
    exact=build_breakdown(b'\n'.join([b'Tanggal,Stok\n2026-10-15,1\n']*5),'exact.csv')
    assert len(tables(exact))==5 and exact['status']=='ready'


@pytest.mark.parametrize('filename,data,status',[
 ('notes.json',b'{"stock":1}','unsupported'),('old.xls',b'opaque','unsupported'),
 ('bad.xlsx',b'not a zip','failed'),('bad.csv',b'\xff\x80','failed'),('nul.csv',b'a,b\n\x00,b\n','failed'),
])
def test_unsupported_or_invalid_is_not_falsely_classified(filename,data,status):
    report=build_breakdown(data,filename)
    assert report['status']==status and report['classification']=='unknown' and not tables(report)


def test_macro_payload_not_processed_and_formulas_not_executed():
    raw=workbook({'Formula':[['Tanggal','Material','Stok'],['2026-10-15','Lateks','=1+2']]})
    report=build_breakdown(raw,'formula.xlsx')
    assert tables(report)[0]['columns'][2]['observed_type']=='empty'
    assert 'Formula/macro tidak dijalankan' in ' '.join(report['notes'])
    stream=io.BytesIO(raw)
    with ZipFile(stream,'a') as zipped:zipped.writestr('xl/vbaProject.bin',b'NO_EXECUTION')
    assert build_breakdown(stream.getvalue(),'macro.xlsx')['status']=='failed'


@pytest.fixture
def client(app_factory):
    app=app_factory(seed_demo=False,start_worker=False,seed_sources=False)
    with TestClient(app) as test_client:yield test_client


def flush(client):
    for _ in range(40):
        if not client.app.state.service.process_next():return
    raise AssertionError('jobs did not drain')


def upload(client,raw=STOCK,name='stok.csv',mode='archive'):
    response=client.post('/api/uploads',headers=A,files={'file':(name,raw)},data={'consent':'true','intake_mode':mode})
    assert response.status_code==202,response.text
    flush(client)
    return response.json()['version']


def get_profile(client,version):
    response=client.get('/api/versions/'+version['id']+'/breakdown',headers=A)
    assert response.status_code==200,response.text
    return response.json()


def review_payload(profile,table_index=0):
    table=tables(profile['profile'])[table_index]
    return {'profile_hash':profile['profile_hash'],'confirmed':True,'tables':[{'table_id':table['id'],'category':table['classification']['category'],
            'columns':[{'column_id':column['id'],'role':column['role'],'datatype':column['observed_type']} for column in table['columns']]}]}


def test_automatic_profile_raw_bytes_and_zero_publication(client):
    version=upload(client);profile=get_profile(client,version)
    assert profile['available'] and profile['can_edit'] and not profile['published']
    assert profile['profile']['classification']=='inventory' and profile['review'] is None
    assert client.get('/api/versions/'+version['id']+'/raw',headers=A).content==STOCK
    v=client.get('/api/versions/'+version['id'],headers=A).json()
    assert v['status']=='archived' and not v['parsed_meta']['sheets']
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==0
    assert client.app.state.service.process_next() is False


def test_purchase_mode_also_gets_profile_without_changing_legacy_mapping(client):
    version=upload(client,mode='purchase');profile=get_profile(client,version)
    assert profile['profile']['classification']=='inventory'
    v=client.get('/api/versions/'+version['id'],headers=A).json()
    assert v['status']=='awaiting_mapping' and v['parsed_meta']['sheets']
    assert client.get('/api/stats?dataset=live',headers=A).json()['observations']==0


def test_profile_and_review_ownership_and_auth(client):
    version=upload(client);profile=get_profile(client,version);base='/api/versions/'+version['id']+'/breakdown'
    for headers,status in (({},401),(B,404)):
        assert client.get(base,headers=headers).status_code==status
        assert client.post(base,headers=headers).status_code==status
        assert client.post(base+'/review',headers=headers,json=review_payload(profile)).status_code==status


def test_review_corrects_metadata_but_preserves_original_suggestions_and_counts(client):
    version=upload(client);profile=get_profile(client,version);body=review_payload(profile)
    body['tables'][0]['category']='unknown';body['tables'][0]['columns'][2].update(role='quantity',datatype='text')
    response=client.post('/api/versions/'+version['id']+'/breakdown/review',headers=A,json=body)
    assert response.status_code==200,response.text
    assert response.json()['raw_changed'] is False and response.json()['published'] is False
    after=get_profile(client,version)
    assert after['profile_hash']==profile['profile_hash'] and after['profile']==profile['profile']
    assert after['review']['tables'][0]['category']=='unknown'
    assert after['review']['tables'][0]['columns'][2]['role']=='quantity'
    assert client.get('/api/versions/'+version['id']+'/raw',headers=A).content==STOCK
    with client.app.state.db.connect() as conn:
        assert conn.execute('SELECT COUNT(*) FROM observations').fetchone()[0]==0
        assert conn.execute('SELECT COUNT(*) FROM excel_templates').fetchone()[0]==0


def test_review_requires_confirmation_current_hash_and_exact_columns(client):
    version=upload(client);profile=get_profile(client,version);base='/api/versions/'+version['id']+'/breakdown/review'
    good=review_payload(profile)
    bad=copy.deepcopy(good);bad['confirmed']=False
    assert client.post(base,headers=A,json=bad).status_code==400
    bad=copy.deepcopy(good);bad['profile_hash']='0'*64
    assert client.post(base,headers=A,json=bad).status_code==409
    mutations=[lambda b:b['tables'][0].update(category='invented'),lambda b:b['tables'][0]['columns'].pop(),
               lambda b:b['tables'][0]['columns'][0].update(role='invented'),lambda b:b['tables'][0]['columns'][0].update(datatype='invented'),
               lambda b:b['tables'][0].update(table_id='unknown')]
    for mutate in mutations:
        bad=copy.deepcopy(good);mutate(bad)
        assert client.post(base,headers=A,json=bad).status_code==400
    bad=copy.deepcopy(good);bad['tables'][0]['columns'][0]['value']='not allowed'
    assert client.post(base,headers=A,json=bad).status_code==422
    assert get_profile(client,version)['review'] is None


def test_review_snapshots_merge_tables_and_idempotent_replay(client):
    raw=workbook({'Stok':[['Tanggal','Stok'],['2026-10-15',1]],'Produksi':[['Tanggal','Hasil produksi'],['2026-10-15',2]]})
    version=upload(client,raw,'two.xlsx');profile=get_profile(client,version);base='/api/versions/'+version['id']+'/breakdown/review'
    one=review_payload(profile,0);two=review_payload(profile,1)
    assert client.post(base,headers=A,json=one).json()['revision']==1
    assert client.post(base,headers=A,json=two).json()['revision']==2
    assert len(get_profile(client,version)['review']['tables'])==2
    replay=client.post(base,headers=A,json=two).json()
    assert replay['reused'] and replay['revision']==2
    with client.app.state.db.connect() as conn:assert conn.execute('SELECT COUNT(*) FROM intake_classification_reviews').fetchone()[0]==2


def test_dedup_and_new_versions_do_not_inherit_review_silently(client):
    one=upload(client);profile=get_profile(client,one);base='/api/versions/'+one['id']+'/breakdown/review'
    assert client.post(base,headers=A,json=review_payload(profile)).status_code==200
    duplicate=upload(client)
    assert duplicate['id']==one['id'] and get_profile(client,duplicate)['review']['revision']==1
    two=upload(client,STOCK.replace(b'100',b'150'))
    assert two['id']!=one['id'] and two['version_no']==2
    assert get_profile(client,two)['review'] is None
    assert client.get('/api/versions/'+one['id']+'/raw',headers=A).content==STOCK


def test_legacy_archive_profile_is_explicitly_queued_without_changing_status(client):
    version=upload(client)
    with client.app.state.db.connect() as conn:
        row=conn.execute('SELECT parsed_meta FROM document_versions WHERE id=?',(version['id'],)).fetchone()
        meta=json.loads(row['parsed_meta']);del meta['breakdown']
        conn.execute('UPDATE document_versions SET parsed_meta=? WHERE id=?',(json.dumps(meta),version['id']))
    base='/api/versions/'+version['id']+'/breakdown'
    assert get_profile(client,version)['available'] is False
    assert client.post(base,headers=A).json()['queued'] is True
    assert client.post(base,headers=A).json()['reused'] is True
    flush(client)
    assert get_profile(client,version)['profile']['classification']=='inventory'
    assert client.get('/api/versions/'+version['id'],headers=A).json()['status']=='archived'
    assert client.post(base,headers=A).json()=={'reused':True,'queued':False}


def test_unsupported_format_and_worker_failure_keep_raw(client,monkeypatch):
    opaque=upload(client,b'RAW opaque','raw.bin')
    assert get_profile(client,opaque)['profile']['status']=='unsupported'
    monkeypatch.setattr('backend.archive.isolated_extract',lambda *args,**kwargs:('','extraction_failed','safe',[],None))
    version=upload(client,STOCK,'worker-budget.csv')
    assert get_profile(client,version)['profile']['status']=='failed'
    assert client.get('/api/versions/'+version['id']+'/raw',headers=A).content==STOCK
    assert client.get('/api/versions/'+version['id'],headers=A).json()['status']=='archived'


def test_delete_cascades_reviews_and_profile(client):
    version=upload(client);profile=get_profile(client,version)
    assert client.post('/api/versions/'+version['id']+'/breakdown/review',headers=A,json=review_payload(profile)).status_code==200
    assert client.delete('/api/documents/'+version['document_id'],headers=A).status_code==200
    with client.app.state.db.connect() as conn:assert conn.execute('SELECT COUNT(*) FROM intake_classification_reviews').fetchone()[0]==0
    assert client.get('/api/versions/'+version['id']+'/breakdown',headers=A).status_code==404


def test_public_profile_readonly_for_private_workspaces(client):
    archived=client.app.state.service.archive('external','public_stok.csv',STOCK,internal=True)
    flush(client);version=archived['version'];base='/api/versions/'+version['id']+'/breakdown'
    profile=client.get(base,headers=B)
    assert profile.status_code==200 and profile.json()['can_edit'] is False
    assert client.post(base,headers=B).status_code==404
    assert client.post(base+'/review',headers=B,json=review_payload(profile.json())).status_code==404


def test_declared_category_does_not_override_inferred_business_classification(client):
    response=client.post('/api/uploads',headers=A,files={'file':('wrong_hint.csv',STOCK)},
                         data={'consent':'true','intake_mode':'archive','category':'quality'})
    assert response.status_code==202;flush(client);version=response.json()['version']
    assert get_profile(client,version)['profile']['classification']=='inventory'
    assert client.get('/api/versions/'+version['id'],headers=A).json()['intake_meta']['category']=='quality'


def test_csv_event_gets_same_breakdown_and_keeps_exact_body(client):
    response=client.post('/api/intake/events?stream=stok&scope_confirmed=true&rights_confirmed=true',
                         headers={**A,'Content-Type':'text/csv'},content=STOCK)
    assert response.status_code==202,response.text
    flush(client);version=response.json()['version']
    assert get_profile(client,version)['profile']['classification']=='inventory'
    assert client.get('/api/versions/'+version['id']+'/raw',headers=A).content==STOCK
