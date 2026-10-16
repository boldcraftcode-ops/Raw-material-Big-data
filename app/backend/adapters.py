"""Source-specific extraction; no generated prices, guessed units, or confirmed news events."""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from email.utils import parsedate_to_datetime
from urllib.parse import urlencode, urljoin, urlsplit
from zipfile import ZipFile

from bs4 import BeautifulSoup
from defusedxml import ElementTree as ET
from openpyxl import load_workbook

CATALOG = {
    'worldbank': {
        'name': 'World Bank · Pink Sheet', 'category': 'price', 'kind': 'market_price',
        'interval': 86400, 'minimum_interval': 21600, 'config': {},
        'description': 'Benchmark bulanan karet RSS3 dan TSR20, USD/kg. Bukan harga spot lokal.',
        'attribution': 'World Bank, Commodity Price Data (The Pink Sheet).',
        'documentation': 'https://www.worldbank.org/en/research/commodity-markets',
        'terms': 'https://www.worldbank.org/en/about/legal/terms-of-use-for-datasets',
        'license': 'World Bank dataset terms / CC BY 4.0, tunduk pada pengecualian pihak ketiga.',
    },
    'faostat_bulk': {
        'name': 'FAOSTAT · Produksi karet', 'category': 'production', 'kind': 'production_stat',
        'interval': 604800, 'minimum_interval': 86400, 'config': {'year_min': 2000},
        'description': 'Produksi tahunan karet alam menurut wilayah dari bulk resmi QCL; default sejak 2000. Bukan statistik Indonesia saja.',
        'attribution': 'FAO. FAOSTAT: Production — Crops and livestock products (QCL).',
        'documentation': 'https://www.fao.org/faostat/en/#data/QCL',
        'terms': 'https://www.fao.org/contact-us/terms/db-terms-of-use/en/',
        'license': 'CC BY 4.0 dan ketentuan database FAO; tidak menyiratkan dukungan FAO.',
    },
    'bmkg': {
        'name': 'BMKG · Prakiraan cuaca', 'category': 'weather', 'kind': 'weather_forecast',
        'interval': 21600, 'minimum_interval': 21600,
        'config': {'adm4': '12.71.01.1001', 'context': 'Pasar Baru, Medan Kota — lokasi referensi regional, belum dipetakan ke kebun/pemasok.'},
        'description': 'Prakiraan 3-harian pada kode desa/kelurahan yang dipantau. Tidak membuktikan kondisi kebun, banjir, atau gangguan pasok.',
        'attribution': 'BMKG — Badan Meteorologi, Klimatologi, dan Geofisika.',
        'documentation': 'https://data.bmkg.go.id/prakiraan-cuaca/',
        'terms': 'https://data.bmkg.go.id/prakiraan-cuaca/',
        'license': 'Data terbuka BMKG; atribusi BMKG wajib. Batas resmi 60 permintaan/menit/IP.',
    },
    'news_rss': {
        'name': 'Google News RSS · Rantai pasok', 'category': 'supply_chain', 'kind': 'news_report',
        'interval': 3600, 'minimum_interval': 1800, 'config': {'keywords': 'supply exports production shipping floods'},
        'description': 'Discovery metadata berita karet alam: judul, penerbit, tanggal, tautan. Artikel tidak di-crawl; laporan belum diverifikasi.',
        'attribution': 'Google News RSS dan masing-masing penerbit yang tercantum.',
        'documentation': 'https://news.google.com/',
        'terms': 'https://policies.google.com/terms',
        'license': 'Metadata feed untuk discovery; hak cipta penerbit tetap berlaku. Bukan lisensi menyalin isi artikel.',
    },
    'gdelt': {
        'name': 'GDELT · Discovery berita', 'category': 'supply_chain', 'kind': 'news_report',
        'interval': 3600, 'minimum_interval': 1800, 'config': {},
        'description': 'Metadata DOC API untuk pencarian natural rubber, maksimum 250 hasil dalam 24 jam; mungkin tidak lengkap.',
        'attribution': 'GDELT Project dan penerbit artikel.',
        'documentation': 'https://blog.gdeltproject.org/gdelt-doc-2-0-api-debuts/',
        'terms': 'https://www.gdeltproject.org/about.html',
        'license': 'Discovery metadata saja; tidak memperoleh hak menyalin artikel dari penerbit.',
    },
}


def validate_config(adapter, config):
    if adapter not in CATALOG: raise ValueError('Adapter tidak didukung.')
    if not isinstance(config, dict): raise ValueError('Konfigurasi harus objek JSON.')
    allowed = set(CATALOG[adapter]['config'])
    if set(config) - allowed: raise ValueError('Opsi konfigurasi tidak dikenal; URL/domain bebas tidak diizinkan.')
    result = {**CATALOG[adapter]['config'], **config}
    if adapter == 'faostat_bulk':
        y = result['year_min']
        if isinstance(y, bool) or not isinstance(y, int) or not 1961 <= y <= 2100: raise ValueError('Tahun minimum harus 1961–2100.')
    if adapter == 'bmkg':
        if not isinstance(result['adm4'], str) or not re.fullmatch(r'\d{2}\.\d{2}\.\d{2}\.\d{4}', result['adm4']):
            raise ValueError('Kode BMKG harus kode adm4 Kemendagri, contoh 12.71.01.1001; bukan kode BPS.')
        if not isinstance(result['context'], str) or not 8 <= len(result['context']) <= 500:
            raise ValueError('Konteks lokasi cuaca perlu dinyatakan (8–500 karakter).')
    if adapter == 'news_rss':
        if not isinstance(result['keywords'], str) or not re.fullmatch(r'[\w\s-]{0,120}', result['keywords'], flags=re.UNICODE):
            raise ValueError('Kata kunci tambahan hanya kata, spasi, atau tanda hubung; maksimum 120 karakter.')
    return result


def number(value, *, minimum=Decimal('0'), maximum=Decimal('1e15')):
    if isinstance(value, bool) or value is None: raise ValueError('Angka kosong/tidak valid.')
    try: d = Decimal(str(value).strip())
    except InvalidOperation: raise ValueError('Angka tidak dapat dibaca tanpa menebak format.')
    if not d.is_finite() or d < minimum or d > maximum: raise ValueError('Angka di luar rentang aman.')
    return format(d, 'f')


def utc(value):
    if not isinstance(value, str): raise ValueError('Waktu sumber tidak tersedia.')
    dt = datetime.fromisoformat(value.replace('Z', '+00:00'))
    if dt.tzinfo is None: dt = dt.replace(tzinfo=timezone.utc)  # API BMKG documents these fields as UTC.
    return dt.astimezone(timezone.utc).isoformat(timespec='seconds')


def safe_link(value):
    u = urlsplit(value or '')
    return value if u.scheme in ('https', 'http') and u.hostname and not u.username and not u.password else ''


def record(key, kind, title, period, fields, artifact, pointer):
    return {'key': key, 'kind': kind, 'title': title[:1000], 'period': period, 'fields': fields,
            'version_id': artifact['version_id'], 'pointer': pointer}


def worldbank(ctx, config):
    catalog = ctx.fetch('catalog', 'https://www.worldbank.org/en/research/commodity-markets', 'worldbank_catalog.html', max_bytes=4 * 1024 * 1024)
    soup = BeautifulSoup(catalog['body'], 'html.parser')
    urls = [urljoin(catalog['url'], a['href']) for a in soup.select('a[href]') if re.search(r'/CMO-Historical-Data-Monthly\.xlsx(?:\?|$)', a['href'], flags=re.I)]
    if not urls: raise ValueError('Tautan workbook bulanan tidak ditemukan; struktur halaman World Bank mungkin berubah.')
    data = ctx.fetch('monthly-prices', urls[0], 'worldbank_monthly_prices.xlsx', max_bytes=8 * 1024 * 1024)
    records, rejected, statistics = parse_worldbank(data)
    return records, rejected, {'scope': 'RSS3 dan TSR20; benchmark bulanan USD/kg; seluruh periode yang tersedia.', 'complete_within_scope': True, 'discovery_url': catalog['url'], **statistics}


def parse_worldbank(data):
    records = []; rejected = []; missing = 0
    with ZipFile(io.BytesIO(data['body'])) as z:
        if len(z.infolist()) > 2000 or sum(i.file_size for i in z.infolist()) > 60 * 1024 * 1024: raise ValueError('Workbook World Bank melampaui batas parser.')
    wb = load_workbook(io.BytesIO(data['body']), read_only=True, data_only=True, keep_links=False)
    found = set()
    try:
        for ws in wb:
            if re.sub(r'[^a-z]', '', ws.title.lower()) != 'monthlyprices': continue
            rows = ws.iter_rows(values_only=True)
            headers = {}; units = {}; header_row = None
            for rn, row in enumerate(rows, 1):
                if rn > 2000: raise ValueError('Workbook melebihi rentang 2.000 baris yang diizinkan adapter harga.')
                if header_row is None:
                    for col, value in enumerate(row):
                        label = re.sub(r'[^a-z0-9]', '', str(value or '').lower())
                        grade = 'RSS3' if label == 'rubberrss3' else 'TSR20' if label == 'rubbertsr20' else None
                        if grade: headers[col] = grade
                    if headers: header_row = rn
                    continue
                if not row: continue
                period_raw = str(row[0] or '').strip()
                match = re.fullmatch(r'(\d{4})M(0[1-9]|1[0-2])', period_raw)
                if not match:
                    for col in headers:
                        if col < len(row) and row[col] is not None:
                            unit = re.sub(r'\s', '', str(row[col]).lower()).replace('us$', '$')
                            if unit in ('($/kg)', '$/kg', 'usd/kg', '(usd/kg)'): units[col] = 'USD/kg'
                    continue
                period = f'{match[1]}-{match[2]}'
                for col, grade in headers.items():
                    if units.get(col) != 'USD/kg': raise ValueError('Unit seri karet tidak dikenali sebagai USD/kg; tidak dikonversi secara implisit.')
                    raw = row[col] if col < len(row) else None
                    if raw is None or str(raw).strip().lower() in ('', '..', '...', '…', 'n.a.', 'n/a', 'na'):
                        missing += 1; continue
                    pointer = {'sheet': ws.title, 'row': rn, 'column': col + 1, 'grade': grade, 'raw_period': period_raw, 'raw_value': raw}
                    try: price = number(raw, maximum=Decimal('1000'))
                    except ValueError as exc:
                        rejected.append((data['version_id'], pointer, str(exc))); continue
                    fields = {'material': 'Natural rubber', 'grade': grade, 'price': price, 'currency': 'USD', 'unit': 'kg',
                              'price_kind': 'monthly_benchmark', 'frequency': 'monthly', 'period': period,
                              'attribution': CATALOG['worldbank']['attribution'], 'not_local_spot': True}
                    records.append(record(f'{grade}:{period}', 'market_price', f'Karet {grade} · {period}', period, fields, data, pointer)); found.add(grade)
            if headers: break
    finally: wb.close()
    if found != {'RSS3', 'TSR20'}: raise ValueError('Dua seri harga karet yang diharapkan tidak lengkap; hasil baru tidak dipublikasikan.')
    periods = {g: [r['period'] for r in records if r['fields']['grade']==g] for g in sorted(found)}
    return records, rejected, {'source_missing_values': missing, 'series_periods': {g: {'first': min(p), 'last': max(p), 'records': len(p)} for g,p in periods.items()}}


def faostat_bulk(ctx, config):
    catalog = ctx.fetch('catalog', 'https://bulks-faostat.fao.org/production/datasets_E.xml', 'faostat_bulk_catalog.xml', max_bytes=2 * 1024 * 1024)
    root = ET.fromstring(catalog['body'], forbid_dtd=True, forbid_entities=True, forbid_external=True)
    dataset = next((d for d in root if d.findtext('DatasetCode') == 'QCL'), None)
    if dataset is None: raise ValueError('Dataset QCL tidak ditemukan dalam katalog resmi FAOSTAT.')
    url = dataset.findtext('FileLocation')
    if not url: raise ValueError('Lokasi bulk FAOSTAT tidak tersedia.')
    data = ctx.fetch('qcl-bulk', url, 'faostat_qcl.zip', max_bytes=64 * 1024 * 1024)
    rows, rejected, coverage = parse_faostat(data, config, heartbeat=ctx.heartbeat, dataset_updated_at=dataset.findtext('DateUpdate'))
    coverage['dataset_updated_at'] = dataset.findtext('DateUpdate')
    coverage['source_rows_declared'] = dataset.findtext('FileRows')
    return rows, rejected, coverage


def parse_faostat(data, config, heartbeat=None, dataset_updated_at=None):
    records = []; rejected = []; scanned = 0; capped = False; missing = 0
    with ZipFile(io.BytesIO(data['body'])) as z:
        members = [i for i in z.infolist() if i.filename.lower().endswith('.csv') and 'all_data' in i.filename.lower() and 'normalized' in i.filename.lower() and not i.is_dir()]
        if len(members) != 1 or len(z.infolist()) > 20 or members[0].file_size > 1024 * 1024 * 1024:
            raise ValueError('Struktur/ukuran bulk QCL tidak sesuai kebijakan ZIP aman.')
        member = members[0]
        flag_names={};flag_member=None
        for entry in z.infolist():
            if entry.filename.lower().endswith('_flags.csv') and entry.file_size<100000:
                flag_member=entry.filename
                with z.open(entry) as f, io.TextIOWrapper(f,encoding='utf-8-sig',newline='') as ft:
                    for fr in csv.DictReader(ft):
                        cleaned={k.strip():v.strip() for k,v in fr.items() if k and v}
                        if cleaned.get('Flag'):flag_names[cleaned['Flag']]=cleaned.get('Description') or cleaned.get('description')
        if member.file_size / max(1, member.compress_size) > 150: raise ValueError('Rasio kompresi bulk terlalu besar.')
        with z.open(member) as f, io.TextIOWrapper(f, encoding='utf-8-sig', newline='') as text:
            reader = csv.DictReader(text)
            required = {'Area', 'Item', 'Element', 'Year', 'Unit', 'Value'}
            if not required.issubset(set(reader.fieldnames or [])): raise ValueError('Skema CSV FAOSTAT berubah; kolom wajib hilang.')
            for rowno, row in enumerate(reader, 2):
                scanned += 1
                if scanned % 200000 == 0 and heartbeat: heartbeat()
                if scanned > 8_000_000: capped = True; break
                item = (row.get('Item') or '').lower()
                if not ('rubber' in item and ('natural' in item or str(row.get('Item Code', '')) == '836')) or 'synthetic' in item: continue
                if (row.get('Element') or '').strip().lower() != 'production': continue
                try: year = int(row['Year'])
                except ValueError: continue
                if year < config['year_min']: continue
                if not row['Value'].strip():
                    missing += 1; continue
                pointer = {'member': member.filename, 'row': rowno, 'raw_fields': row,'flag_lookup_member':flag_member}
                try:
                    if row['Unit'].strip().lower() not in ('t', 'tonnes', 'tonne'): raise ValueError('Unit produksi bukan ton metrik yang dikenali.')
                    value = number(row['Value'])
                except ValueError as exc:
                    rejected.append((data['version_id'], pointer, str(exc))); continue
                code = row.get('Area Code') or row.get('Area Code (M49)') or row['Area']
                key = f'{code}:rubber:production:{year}'
                fields = {'material': row['Item'], 'area': row['Area'], 'area_code': code, 'year': year,
                          'production': value, 'unit': 'tonnes', 'frequency': 'annual', 'flag': row.get('Flag'),
                          'flag_description':flag_names.get(row.get('Flag')),'dataset_updated_at':dataset_updated_at,'source_url':'https://www.fao.org/faostat/en/#data/QCL',
                          'note': row.get('Note'), 'element_code': row.get('Element Code'),
                          'attribution': 'FAO. '+(str(dataset_updated_at)[:4] if dataset_updated_at else 'tahun pembaruan pada metadata sumber')+'. FAOSTAT: Crops and livestock products (QCL). https://www.fao.org/faostat/en/#data/QCL. Licence: CC-BY-4.0.', 'license': 'CC BY 4.0',
                          'area_scope': 'Wilayah sebagaimana dilaporkan FAOSTAT; agregat regional bukan negara terpisah.'}
                records.append(record(key, 'production_stat', f'Produksi karet · {row["Area"]} · {year}', str(year), fields, data, pointer))
                if len(records) > 20000: raise ValueError('Hasil karet melampaui 20.000 record per run; sempitkan konfigurasi.')
    if not records: raise ValueError('Tidak ada produksi karet alam yang lolos validasi dalam bulk QCL.')
    return records, rejected, {'scope': f'Produksi karet alam tahunan sejak {config["year_min"]}; wilayah dan flag sebagaimana FAOSTAT.',
                               'complete_within_scope': not capped, 'bulk_rows_scanned': scanned,
                               'source_missing_values': missing, 'raw_contains_other_commodities': True, 'note': 'Bulk raw mengandung komoditas lain; hanya produksi karet alam dinormalisasi.'}


def bmkg(ctx, config):
    url = 'https://api.bmkg.go.id/publik/prakiraan-cuaca?' + urlencode({'adm4': config['adm4']})
    data = ctx.fetch('forecast:' + config['adm4'], url, 'bmkg_' + config['adm4'] + '.json', max_bytes=2 * 1024 * 1024)
    return parse_bmkg(data, config)


def parse_bmkg(data, config):
    obj = json.loads(data['body']); location = obj.get('lokasi') or {}
    if location.get('adm4') != config['adm4']: raise ValueError('Lokasi respons BMKG tidak sama dengan kode yang dipantau.')
    records = []; rejected = []
    for block_no, block in enumerate(obj.get('data') or []):
        loc = block.get('lokasi') or location
        if loc.get('adm4') != config['adm4']: continue
        for day_no, day in enumerate(block.get('cuaca') or []):
            for index, forecast in enumerate(day):
                pointer = {'json_pointer': f'/data/{block_no}/cuaca/{day_no}/{index}', 'raw_fields': forecast}
                try:
                    valid = utc(forecast['utc_datetime'])
                    issued = utc(forecast['analysis_date']) if forecast.get('analysis_date') else None
                    temp = number(forecast.get('t'), minimum=Decimal('-60'), maximum=Decimal('70'))
                    humidity = number(forecast.get('hu'), maximum=Decimal('100'))
                    wind = number(forecast.get('ws'), maximum=Decimal('500'))
                except (ValueError, KeyError, TypeError) as exc:
                    rejected.append((data['version_id'], pointer, 'Waktu/nilai prakiraan BMKG tidak valid.')); continue
                fields = {'material_context': 'Natural rubber monitoring context; not a plantation observation',
                          'location': location, 'monitoring_context': config['context'], 'valid_time_utc': valid,
                          'issued_at_utc': issued, 'local_datetime_reported': forecast.get('local_datetime'),
                          'temperature_c': temp, 'humidity_percent': humidity, 'wind_kmh': wind,
                          'weather': forecast.get('weather_desc'), 'weather_code': forecast.get('weather'),
                          'data_kind': 'forecast_not_observation', 'attribution': CATALOG['bmkg']['attribution']}
                title = f'{location.get("desa", config["adm4"])} · {forecast.get("weather_desc") or "Prakiraan"} · {valid[:16]}'
                records.append(record(config['adm4'] + ':' + valid, 'weather_forecast', title, valid, fields, data, pointer))
    if not records: raise ValueError('Respons BMKG tidak berisi interval prakiraan yang valid.')
    return records, rejected, {'scope': 'Satu kode adm4: ' + config['adm4'], 'complete_within_scope': True,
                               'location': location, 'context': config['context'], 'forecast_not_observation': True,
                               'intervals_received':len(records),'valid_first':min(r['period'] for r in records),'valid_last':max(r['period'] for r in records),
                               'horizon_note':'Cakupan mengikuti interval yang dikembalikan API; tidak menjamin 72 jam penuh.',
                               'plantation_link_verified': False, 'attribution': CATALOG['bmkg']['attribution']}


def news_rss(ctx, config):
    words = config['keywords'].split()
    query = '("natural rubber" OR "karet alam" OR bokar)' + (' (' + ' OR '.join(words) + ')' if words else '') + ' when:7d'
    url = 'https://news.google.com/rss/search?' + urlencode({'q': query, 'hl': 'en-US', 'gl': 'US', 'ceid': 'US:en'})
    data = ctx.fetch('news-feed', url, 'rubber_supply_news.rss', max_bytes=4 * 1024 * 1024)
    root = ET.fromstring(data['body'], forbid_dtd=True, forbid_entities=True, forbid_external=True)
    channel = root.find('channel')
    if channel is None: raise ValueError('Respons bukan RSS berita yang dikenali.')
    items = channel.findall('item')
    records = []; rejected = []
    for index, item in enumerate(items[:200]):
        title = BeautifulSoup(item.findtext('title') or '', 'html.parser').get_text(' ', strip=True)
        link = safe_link(item.findtext('link')); guid = item.findtext('guid') or link
        pointer = {'xpath': f'/rss/channel/item[{index + 1}]'}
        if not title or not link or not guid:
            rejected.append((data['version_id'], pointer, 'Metadata berita tidak mempunyai judul/tautan/identifier yang valid.')); continue
        try: published = parsedate_to_datetime(item.findtext('pubDate') or '').astimezone(timezone.utc).isoformat(timespec='seconds')
        except (ValueError, TypeError): published = None
        source = item.find('source')
        publisher = source.text if source is not None else None
        publisher_url = safe_link(source.get('url')) if source is not None else ''
        fields = {'title': title, 'url': link, 'publisher': publisher, 'publisher_url': publisher_url,
                  'published_at': published, 'confirmation_status': 'unverified_report', 'relevance': 'query_match_unreviewed',
                  'event_type': 'supply_chain_news_candidate', 'full_article_collected': False,
                  'attribution': 'Google News RSS; ' + (publisher or 'penerbit pada tautan sumber')}
        records.append(record(guid, 'news_report', title, published or '', fields, data, pointer))
    return records, rejected, {'scope': 'Metadata RSS untuk query karet alam, jendela permintaan 7 hari; bukan seluruh berita dunia.',
                               'query': query, 'complete_within_scope': False, 'feed_items': len(items),
                               'limited_feed': True, 'full_articles_collected': False, 'feed_built_at': channel.findtext('lastBuildDate')}


def gdelt(ctx, config):
    url = 'https://api.gdeltproject.org/api/v2/doc/doc?' + urlencode({'query': '"natural rubber"', 'mode': 'ArtList', 'format': 'json', 'maxrecords': 250, 'timespan': '24h', 'sort': 'DateDesc'})
    data = ctx.fetch('articles', url, 'gdelt_rubber.json', max_bytes=4 * 1024 * 1024)
    obj = json.loads(data['body'])
    if not isinstance(obj, dict) or not isinstance(obj.get('articles'), list): raise ValueError('Skema GDELT tidak dikenali; mungkin pesan batas akses.')
    records = []; rejected = []
    for i, a in enumerate(obj['articles'][:250]):
        link = safe_link(a.get('url')); title = a.get('title') or ''
        pointer = {'json_pointer': f'/articles/{i}'}
        if not link or not title:
            rejected.append((data['version_id'], pointer, 'Artikel GDELT tidak mempunyai metadata wajib.')); continue
        fields = {'title': title, 'url': link, 'publisher': a.get('domain'), 'published_at': None,
                  'seen_at_reported': a.get('seendate'), 'source_country_reported': a.get('sourcecountry'),
                  'language': a.get('language'), 'confirmation_status': 'unverified_report', 'relevance': 'query_match_unreviewed',
                  'full_article_collected': False, 'attribution': 'GDELT Project dan ' + (a.get('domain') or 'penerbit artikel')}
        records.append(record(link, 'news_report', title, a.get('seendate') or '', fields, data, pointer))
    return records, rejected, {'scope': 'GDELT DOC ArtList, natural rubber, 24h, maksimum 250 metadata artikel.',
                               'complete_within_scope': False, 'possibly_truncated': len(obj['articles']) >= 250, 'full_articles_collected': False}


ADAPTERS = {'worldbank': worldbank, 'faostat_bulk': faostat_bulk, 'bmkg': bmkg, 'news_rss': news_rss, 'gdelt': gdelt}
