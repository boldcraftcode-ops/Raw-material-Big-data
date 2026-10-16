from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from .excel_assistant import HEADER_SCAN, choose_header

MAX_ROWS=10000
MAX_COLUMNS=64
MAX_SHEETS=10
MAX_CELL=4096


def serial_cell(value):
    if isinstance(value,(datetime,date)): return value.isoformat()
    if value is None: return None
    if isinstance(value,(str,int,float,bool)):
        if isinstance(value,str) and len(value)>MAX_CELL: raise ValueError('Ada sel melebihi batas 4.096 karakter.')
        return value
    text=str(value)
    if len(text)>MAX_CELL: raise ValueError('Ada sel melebihi batas karakter.')
    return text


def columns_for(values):
    if len(values)>MAX_COLUMNS: raise ValueError('Maksimum 64 kolom per lembar pada MVP ini.')
    if not values or not any(v is not None and str(v).strip() for v in values): raise ValueError('Header kolom tidak ditemukan.')
    return [{'key':f'c{i}','label':str(v).strip() if v is not None and str(v).strip() else f'Kolom {i+1}'} for i,v in enumerate(values)]


def parse_file(path: Path, filename: str, header_overrides=None):
    """Source is read-only; all row pointers keep physical source coordinates."""
    suffix=Path(filename).suffix.lower()
    sheets, rows, warnings=[], [], []
    header_overrides=header_overrides or {}

    def add_sheet(name, records):
        if not any(any(value is not None and str(value).strip() for value in r['values']) for r in records): return
        selected, candidates=choose_header(records,header_overrides.get(name))
        columns=selected['columns']
        if len(columns)>MAX_COLUMNS: raise ValueError('Maksimum 64 kolom per lembar pada MVP ini.')
        count=0
        for record in records[selected['index']+selected['header_depth']:]:
            values=record['values']
            if not any(value is not None and str(value).strip() for value in values): continue
            if any(value is not None and str(value).strip() for value in values[len(columns):]):
                raise ValueError(f"Lembar {name}, baris {record['row_no']}: ada data di luar header. Periksa pilihan header atau arsipkan sebagai raw.")
            cells={column['key']:values[i] if i<len(values) else None for i,column in enumerate(columns)}
            rows.append({'sheet':name,'row_no':record['row_no'],'line_end':record['line_end'],'cells':cells})
            count+=1
            if len(rows)>MAX_ROWS: raise ValueError('Maksimum total 10.000 baris data per file.')
        sheets.append({'name':name,'columns':columns,'row_count':count,'suggested_mapping':selected['suggested_mapping'],
                       'header_row':selected['header_row'],'header_depth':selected['header_depth'],
                       'header_end_row':selected['header_end_row'],'header_score':selected['score'],
                       'header_candidates':candidates})
        if selected['header_row']>1 or selected['header_depth']>1:
            warnings.append(f"{name}: header dipilih pada baris {selected['header_row']}–{selected['header_end_row']}. Bagian sebelumnya tidak dinormalisasi, tetapi tetap ada dalam file asli. Periksa pilihan ini.")

    if suffix=='.csv':
        data=path.read_bytes()
        if b'\x00' in data[:200] and not data.startswith((b'\xff\xfe',b'\xfe\xff')): raise ValueError('CSV tampak seperti berkas biner; unggah CSV teks.')
        encoding='utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig'
        try: text=data.decode(encoding)
        except UnicodeDecodeError: raise ValueError('CSV harus UTF-8 atau UTF-16 dengan BOM. Simpan ulang file menggunakan UTF-8.')
        dialect=csv.excel
        physical_lines=text.splitlines()
        # Leading titles often defeat Sniffer. A later bounded slice may reveal the table delimiter.
        for offset in range(min(HEADER_SCAN,len(physical_lines))):
            try:
                dialect=csv.Sniffer().sniff('\n'.join(physical_lines[offset:offset+20])[:12000],delimiters=',;\t|')
                break
            except csv.Error: pass
        reader=csv.reader(io.StringIO(text,newline=''),dialect=dialect)
        records=[]
        previous=0
        for values in reader:
            start=previous+1; previous=reader.line_num
            if len(values)>MAX_COLUMNS: raise ValueError('Maksimum 64 kolom per lembar pada MVP ini.')
            records.append({'row_no':start,'line_end':reader.line_num,'values':[serial_cell(value) for value in values]})
            if len(records)>MAX_ROWS+HEADER_SCAN+2: raise ValueError('Rentang baris file terlalu besar; raw tetap diarsipkan.')
        if not records: raise ValueError('File CSV kosong.')
        add_sheet('CSV',records)
    elif suffix=='.xlsx':
        try:
            with ZipFile(path) as archive:
                entries=archive.infolist()
                if len(entries)>2000 or sum(item.file_size for item in entries)>60*1024*1024: raise ValueError('Ukuran isi XLSX terlalu besar untuk diproses dengan aman.')
                names=[item.filename.lower() for item in entries]
                if any('vbaproject' in name for name in names): raise ValueError('Workbook dengan macro tidak didukung.')
                if '[content_types].xml' not in names or not any(name.startswith('xl/') for name in names): raise ValueError('File bukan workbook XLSX yang valid.')
        except BadZipFile: raise ValueError('Struktur XLSX tidak valid.')
        handle=path.open('rb')
        try: workbook=load_workbook(handle,read_only=True,data_only=True,keep_links=False)
        except Exception:
            handle.close()
            raise ValueError('Workbook XLSX tidak dapat dibaca. Simpan ulang sebagai XLSX standar.')
        try:
            visible=[sheet for sheet in workbook if sheet.sheet_state=='visible']
            if len(visible)>MAX_SHEETS: raise ValueError('Maksimum 10 lembar terlihat pada MVP ini.')
            if any(sheet.sheet_state!='visible' for sheet in workbook): warnings.append('Lembar tersembunyi tidak diproses; file asli tetap diarsipkan.')
            warnings.append('Formula tidak dijalankan. Sel formula hanya dibaca dari nilai cache; nilai yang tidak tersedia memerlukan review.')
            for sheet in visible:
                if sheet.max_row and sheet.max_row>MAX_ROWS+100: raise ValueError('Rentang baris lembar terlalu besar. Ekspor rentang data saja ke CSV/XLSX baru.')
                if sheet.max_column and sheet.max_column>256: raise ValueError('Rentang kolom lembar terlalu besar; raw tetap tersimpan. Ekspor rentang tabel yang diperlukan.')
                records=[]
                for row_no, values in enumerate(sheet.iter_rows(values_only=True),start=1):
                    values=list(values)
                    while values and values[-1] is None: values.pop()
                    if len(values)>MAX_COLUMNS: raise ValueError('Maksimum 64 kolom per lembar pada MVP ini.')
                    records.append({'row_no':row_no,'line_end':row_no,'values':[serial_cell(value) for value in values]})
                add_sheet(sheet.title,records)
        finally:
            workbook.close(); handle.close()
        if len(sheets)>1: warnings.append('Pilih satu lembar saat pemetaan. Lembar lain tidak otomatis dipublikasikan sebagai observasi.')
    else: raise ValueError('Pemetaan tabel mendukung CSV/XLSX; gunakan arsip umum untuk format lain.')
    if not rows: raise ValueError('Header ditemukan, tetapi tidak ada baris data. Isi template sebelum mengunggah.')
    if not sheets: raise ValueError('Tidak ada lembar data yang dapat dibaca.')
    json.dumps(rows,allow_nan=False)
    return {'sheets':sheets,'warnings':warnings,'parser_version':'tabular-v2-local-assistant','row_count':len(rows)},rows
