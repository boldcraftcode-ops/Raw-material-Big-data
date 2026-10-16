"""General raw tabular profiling. Deterministic suggestions, never purchase publication."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from .domain import label_key
from .excel_assistant import ALIASES, header_columns

VERSION='intake-breakdown-v1'
MAX_ROWS=5000
MAX_COLUMNS=64
MAX_SHEETS=10
MAX_TABLES=20
CATEGORIES={'purchase':'Pembelian','sales':'Penjualan','price':'Harga / benchmark','inventory':'Stok / persediaan',
            'production':'Produksi','quality':'Mutu / hasil uji','logistics':'Logistik / pengiriman',
            'weather':'Cuaca','traceability':'Ketertelusuran','mixed':'Campuran','unknown':'Belum terklasifikasi'}
EXTRA_ROLES={
 'date':['tgl masuk','tanggal produksi','tanggal uji','tanggal kirim','tanggal stok','tanggal penjualan','periode','bulan','tahun'],
 'buyer':['pembeli','pelanggan','customer','buyer','nama pembeli','nama pelanggan'],
 'quantity':['jumlah','qty','kuantitas','berat','volume','berat basah','berat kering','berat basah kg','berat kering kg'],
 'stock':['stok','stock','saldo','stok akhir','stok awal','saldo akhir','saldo awal','closing stock','opening stock','persediaan','saldo kg','stok kg'],
 'stock_in':['barang masuk','masuk kg','stok masuk','penerimaan kg'],
 'stock_out':['barang keluar','keluar kg','stok keluar','pengeluaran kg'],
 'warehouse':['gudang','warehouse','kode gudang'],
 'production_output':['hasil produksi','produksi','production','output','produksi kg','output kg','jumlah produksi','hasil olah','production output'],
 'production_line':['mesin','line produksi','lini produksi','production line'],
 'quality_metric':['kadar abu','kadar kotoran','viskositas','pri','po','ash content','dirt content','volatile matter','parameter mutu'],
 'test_result':['hasil uji','nilai uji','test result'],
 'sample_id':['id sampel','nomor sampel','no sampel','sample id'],
 'shipment_id':['no surat jalan','nomor surat jalan','surat jalan','no resi','nomor resi','no kontainer','container','resi'],
 'destination':['tujuan','destination','lokasi tujuan'],
 'origin':['asal','origin','asal bahan','asal kebun','lokasi asal'],
 'lot':['lot','no lot','nomor lot','batch','no batch','kode batch','batch id'],
 'latitude':['latitude','lat','lintang'], 'longitude':['longitude','lon','lng','bujur'],
 'rainfall':['curah hujan','rainfall','curah hujan mm'],
 'temperature':['suhu','temperatur','temperature','suhu c'],
 'humidity':['kelembapan','humidity','kelembaban'],
}
ROLES={key:set(values) for key,values in ALIASES.items()}
for key,values in EXTRA_ROLES.items():ROLES.setdefault(key,set()).update(label_key(value) for value in values)
ROLES['unknown']=set()
INVERSE={}
for role,aliases in ROLES.items():
    for alias in aliases:INVERSE.setdefault(alias,[]).append(role)


def _present(value):return value is not None and str(value).strip()!=''


def display(value,limit=120):
    if value is None:return None
    if isinstance(value,(datetime,date)):return value.isoformat()
    if isinstance(value,float) and not math.isfinite(value):return str(value)
    if isinstance(value,(int,float,bool)):return value
    return str(value)[:limit]


def primitive(value):
    if not _present(value):return 'empty'
    if isinstance(value,bool):return 'boolean'
    if isinstance(value,(datetime,date)):return 'date'
    if isinstance(value,(int,float)):
        return 'number' if not isinstance(value,float) or math.isfinite(value) else 'invalid_number'
    text=str(value).strip()
    if re.match(r'^\d{4}-\d{2}-\d{2}(?:T.*)?$',text):return 'date_text_iso'
    if re.fullmatch(r'\d{1,2}[/\-]\d{1,2}[/\-]\d{4}',text):return 'date_text_ambiguous'
    if re.fullmatch(r'\d{4}-\d{2}',text):return 'period_text'
    if re.fullmatch(r'[+-]?[\d.,]+\s*%',text):return 'percent_text'
    if re.fullmatch(r'(?:Rp\.?|IDR|USD|US\$|\$)\s*[+-]?[\d.,]+',text,re.I):return 'money_text'
    if re.fullmatch(r'[+-]?[\d.,]+',text):
        if re.fullmatch(r'0\d+',text):return 'code_text'
        return 'number_text'
    if text.startswith(('#REF!','#DIV/0!','#VALUE!','#N/A')):return 'excel_error'
    return 'text'


def role_for(label):
    parts=[part.strip() for part in label.split(' / ')]
    for value in [parts[-1],label]+list(reversed(parts[:-1])):
        matches=INVERSE.get(label_key(value),[])
        if len(matches)==1:return matches[0],f'Alias header terdaftar: {value[:100]}.',[]
        if len(matches)>1:return 'unknown','Header cocok dengan lebih dari satu peran; perlu review.',matches
    return 'unknown','Belum ada alias header yang cukup jelas; tidak ditebak dari besar angka.',[]


def classify(columns):
    roles={column['role'] for column in columns}
    matches=[]
    def add(category,condition,evidence):
        if condition:matches.append({'category':category,'roles':sorted(roles & set(evidence))})
    add('purchase',{'supplier','price'}<=roles and bool(roles & {'date','material'}),['supplier','price','date','material'])
    add('sales',{'buyer','price'}<=roles and bool(roles & {'date','material'}),['buyer','price','date','material'])
    add('price','price' in roles and bool(roles & {'date','material'}) and not roles & {'supplier','buyer'},['price','date','material'])
    add('inventory',('stock' in roles or {'stock_in','stock_out'}<=roles) and bool(roles & {'material','warehouse','date'}),['stock','stock_in','stock_out','material','warehouse','date'])
    add('production','production_output' in roles and bool(roles & {'material','production_line','date'}),['production_output','material','production_line','date'])
    add('quality',('quality_metric' in roles or {'test_result','sample_id'}<=roles) and bool(roles & {'material','date','sample_id'}),['quality_metric','test_result','sample_id','material','date'])
    add('logistics','shipment_id' in roles and bool(roles & {'destination','origin','date'}),['shipment_id','destination','origin','date'])
    add('weather',len(roles & {'rainfall','temperature','humidity'})>=2,['rainfall','temperature','humidity','date','location'])
    add('traceability','lot' in roles and bool(roles & {'origin','latitude','longitude'}),['lot','origin','latitude','longitude'])
    category=matches[0]['category'] if len(matches)==1 else 'mixed' if matches else 'unknown'
    return {'category':category,'matches':matches,'reason':'Kombinasi peran header yang terdaftar; saran, bukan verifikasi isi atau keputusan bisnis.' if matches else 'Indikator belum cukup untuk menentukan jenis bisnis.'}


def aggregate(categories):
    categories=set(categories)
    return next(iter(categories)) if len(categories)==1 else 'mixed' if categories else 'unknown'


def _letter(index):
    result=''
    while index:
        index,rest=divmod(index-1,26);result=chr(65+rest)+result
    return result


def _candidate(records,index,initial=False):
    first=records[index]['values']
    if not any(_present(value) for value in first):return None
    if any(primitive(value) not in ('empty','text','code_text') for value in first):return None
    choices=[]
    for depth in (1,2):
        second=records[index+1]['values'] if depth==2 and index+1<len(records) else None
        if depth==2 and not any(role_for(str(value))[0]!='unknown' for value in first if _present(value)):continue
        if depth==2 and (not second or any(primitive(value) not in ('empty','text','code_text') for value in second)):continue
        columns=header_columns(first,second)
        known=sum(role_for(column['label'])[0]!='unknown' for column in columns)
        nonempty=sum(not column['source_blank'] for column in columns)
        if not nonempty:continue
        if not initial and known<2:continue
        following=records[index+depth:index+depth+3]
        if not following:continue
        width=len(columns)
        good=sum(any(_present(value) for value in row['values']) and not any(_present(value) for value in row['values'][width:]) for row in following)
        if not good:continue
        score=known*5+nonempty*.08+good*.3-(depth-1)*.5
        choices.append({'index':index,'depth':depth,'columns':columns,'known':known,'score':score})
    return max(choices,key=lambda choice:choice['score']) if choices else None


def _table(sheet,records,header,end,complete,span):
    start=header['index']+header['depth']; columns=header['columns']
    body=[row for row in records[start:end] if any(_present(value) for value in row['values'][span[0]:span[1]])]
    selected=[]
    for offset in range(span[0],span[1]):
        source=columns[offset]
        role,reason,conflicts=role_for(source['label'])
        values=[(row,row['values'][offset] if offset<len(row['values']) else None) for row in body]
        counts=Counter(primitive(value) for _,value in values)
        kinds={kind for kind,count in counts.items() if count and kind!='empty'}
        dtype=next(iter(kinds)) if len(kinds)==1 else 'mixed' if kinds else 'empty'
        samples=[]
        for row,value in values:
            if _present(value) and len(samples)<2:samples.append({'row_no':row['row_no'],'line_end':row['line_end'],'value':display(value,80)})
        hints=[]
        label=source['label']
        for regex,hint in [(r'\bkg\b','kg'),(r'\b(?:ton|tonne)\b','t'),(r'\b(?:rp|idr)\b','IDR'),(r'\b(?:usd|us\$)\b','USD'),(r'%','%')]:
            if re.search(regex,label,re.I):hints.append(hint)
        selected.append({'id':f'c{offset}','source_column':offset+1,'name':label[:160],
                         'role':role,'role_reason':reason,'role_conflicts':conflicts,'observed_type':dtype,
                         'type_counts':dict(counts),'empty_count':counts.get('empty',0),'nonempty_count':len(body)-counts.get('empty',0),
                         'samples':samples,'unit_hints':hints,'source_ref':f'{sheet}!{_letter(offset+1)}{records[header["index"]]["row_no"]}'})
    header_row=records[header['index']]['row_no'];header_end=records[start-1]['line_end']
    key=hashlib.sha256(f'{VERSION}|{sheet}|{header_row}|{header_end}|{span}'.encode()).hexdigest()[:24]
    preview=[{'row_no':row['row_no'],'line_end':row['line_end'],'cells':{f'c{i}':display(row['values'][i] if i<len(row['values']) else None) for i in range(span[0],min(span[1],span[0]+12))}} for row in body[:4]]
    return {'id':key,'sheet':sheet,'header_row':header_row,'header_end_row':header_end,'header_depth':header['depth'],
            'start_column':span[0]+1,'end_column':span[1],'data_start_row':body[0]['row_no'] if body else None,
            'data_end_row':body[-1]['line_end'] if body else None,'record_count':len(body),'scan_complete':complete,
            'header_detection':'rules' if header['known']>=2 else 'structural_guess','columns':selected,
            'classification':classify(selected),'preview':preview,
            'notes':['Hitungan dan tipe hanya untuk baris yang dipindai; peran/jenis bisnis masih saran.','Preview dibatasi 4 baris / 12 kolom; nilai tampilan dipotong, raw tidak dipotong.']}


def _profile_sheet(name,records,complete,declared_rows,remaining):
    tables=[];assigned=set();covered={};segments=[];start=0;limited=False
    for i,row in enumerate(records):
        if not any(_present(value) for value in row['values']):
            if start<i:segments.append((start,i))
            start=i+1
    if start<len(records):segments.append((start,len(records)))
    for begin,end in segments:
        section=records[:end]
        choices=[choice for i in range(begin,min(end,begin+50)) if (choice:=_candidate(section,i,True))]
        strong=[choice for choice in choices if choice['known']>=2]
        chosen=min(strong,key=lambda choice:choice['index']) if strong else max(choices,key=lambda choice:choice['score']) if choices else None
        while chosen and len(tables)<min(5,remaining):
            body_start=chosen['index']+chosen['depth']
            following=next((choice for i in range(body_start+1,end) if (choice:=_candidate(section,i)) and choice['known']>=2),None)
            stop=following['index'] if following else end
            width=len(chosen['columns'])
            active=[any(_present(row['values'][col]) for row in records[chosen['index']:stop] if col<len(row['values'])) for col in range(width)]
            runs=[];run=None
            for col,used in enumerate(active+[False]):
                if used and run is None:run=col
                if not used and run is not None:runs.append((run,col));run=None
            # Only split side-by-side groups when an entirely blank separator and known headers support both.
            spans=runs if len(runs)>1 and all(sum(role_for(c['label'])[0]!='unknown' for c in chosen['columns'][lo:hi])>=2 for lo,hi in runs) else [(0,width)]
            for span in spans:
                if len(tables)>=min(5,remaining):limited=True;break
                table=_table(name,records,chosen,stop,complete,span)
                if table['record_count']:
                    tables.append(table)
                    assigned.update(range(chosen['index'],stop))
                    for row_index in range(chosen['index'],stop):covered.setdefault(row_index,set()).update(range(*span))
            chosen=following
        if chosen and len(tables)>=min(5,remaining):limited=True
    unassigned=[row['row_no'] for i,row in enumerate(records) if i not in assigned and any(_present(value) for value in row['values'])]
    unassigned_cells=sum(_present(value) and column not in covered.get(i,set()) for i,row in enumerate(records) for column,value in enumerate(row['values']))
    return {'name':name,'declared_rows':declared_rows,'rows_scanned':records[-1]['line_end'] if records else 0,'records_scanned':len(records),'scan_complete':complete,
            'tables':tables,'layout_limited':limited,'unassigned_nonempty_cells':unassigned_cells,'unassigned_nonempty_rows':len(unassigned),'unassigned_row_examples':unassigned[:12],
            'classification':aggregate(table['classification']['category'] for table in tables),
            'notes':['Bagian yang tidak menjadi tabel tetap ada pada raw. Tabel/crosstab rumit mungkin perlu peninjauan struktur.']+(['Batas jumlah tabel dapat membatasi breakdown; periksa cakupan.'] if limited else [])}


def build_breakdown(data,filename):
    suffix=Path(filename).suffix.lower()
    report={'schema_version':VERSION,'mode':'local_rules','external_ai_used':False,'published_as_transactions':False,
            'format':suffix.lstrip('.') or 'unknown','status':'ready','sheets':[],'classification':'unknown',
            'summary':{'sheets_total':0,'sheets_scanned':0,'tables_detected':0,'physical_rows_scanned':0,'data_records_scanned':0,'unassigned_nonempty_rows':0,'unassigned_nonempty_cells':0},
            'notes':[],'limits':{'rows':MAX_ROWS,'columns':MAX_COLUMNS,'sheets':MAX_SHEETS,'tables':MAX_TABLES}}
    if suffix not in ('.csv','.tsv','.xlsx'):
        report.update(status='unsupported',notes=['Breakdown MVP hanya untuk CSV/TSV/XLSX. Format lain tetap raw; tidak diklasifikasikan dari tebakan isi.']);return report
    if len(data)>10*1024*1024:
        report.update(status='partial',notes=['Breakdown dibatasi file 10 MB; raw yang diterima melalui batas internal tetap utuh.']);return report
    try:
        left=MAX_ROWS
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            with ZipFile(io.BytesIO(data)) as archive:
                info=archive.infolist()
                if len(info)>2000 or sum(item.file_size for item in info)>60*1024*1024:raise ValueError('XLSX safety bound')
                if any('vbaproject' in item.filename.lower() for item in info):raise ValueError('Macro not processed')
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
            try:
                report['summary']['sheets_total']=len(book.sheetnames)
                for sheet in book:
                    if sheet.sheet_state!='visible':
                        report['status']='partial';report['notes'].append(f'Lembar tersembunyi tidak dipindai: {sheet.title[:80]}.');continue
                    if len(report['sheets'])>=MAX_SHEETS or left<=0:
                        report['status']='partial';report['notes'].append('Batas sheet/baris total tercapai.');break
                    records=[];complete=True
                    if (sheet.max_column or 0)>MAX_COLUMNS:complete=False
                    for row_no,values in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS+1,values_only=True),1):
                        if len(records)>=left:complete=False;break
                        if any(_present(value) for value in values[MAX_COLUMNS:]):complete=False
                        values=list(values[:MAX_COLUMNS])
                        while values and not _present(values[-1]):values.pop()
                        records.append({'row_no':row_no,'line_end':row_no,'values':values})
                    left-=len(records)
                    report['sheets'].append(_profile_sheet(sheet.title,records,complete,sheet.max_row,MAX_TABLES-sum(len(item['tables']) for item in report['sheets'])))
                    if not complete:report['status']='partial'
            finally:book.close()
            report['notes'].append('Formula/macro tidak dijalankan; hanya nilai cache. Cache formula kosong terlihat sebagai kosong, bukan nilai nol.')
        else:
            encoding='utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig'
            text=data.decode(encoding)
            if '\x00' in text:raise ValueError('Binary')
            lines=text[:64000].splitlines();dialect=csv.excel_tab if suffix=='.tsv' else csv.excel
            for offset in range(min(50,len(lines))):
                try:dialect=csv.Sniffer().sniff('\n'.join(lines[offset:offset+20])[:12000],delimiters=',;\t|');break
                except csv.Error:pass
            reader=csv.reader(io.StringIO(text,newline=''),dialect);records=[];previous=0;complete=True
            for values in reader:
                if len(records)>=MAX_ROWS:complete=False;break
                if len(values)>MAX_COLUMNS:complete=False
                records.append({'row_no':previous+1,'line_end':reader.line_num,'values':values[:MAX_COLUMNS]});previous=reader.line_num
            report['summary']['sheets_total']=1
            report['sheets'].append(_profile_sheet('CSV' if suffix=='.csv' else 'TSV',records,complete,None,MAX_TABLES))
            if not complete:report['status']='partial'
        tables=[table for sheet in report['sheets'] for table in sheet['tables']]
        report['summary'].update(sheets_scanned=len(report['sheets']),tables_detected=len(tables),
                                 physical_rows_scanned=sum(sheet['rows_scanned'] for sheet in report['sheets']),
                                 data_records_scanned=sum(table['record_count'] for table in tables),
                                 unassigned_nonempty_rows=sum(sheet['unassigned_nonempty_rows'] for sheet in report['sheets']),
                                 unassigned_nonempty_cells=sum(sheet['unassigned_nonempty_cells'] for sheet in report['sheets']))
        report['classification']=aggregate(table['classification']['category'] for table in tables)
        if any(not sheet['scan_complete'] or sheet.get('layout_limited') for sheet in report['sheets']):report['status']='partial'
        report['notes']+=['Batas 5.000 baris berarti baris XLSX atau record logis CSV/TSV; baris fisik multiline tetap ditunjuk lewat rentang baris.',
                         'Jenis bisnis dan peran kolom adalah saran dari aturan/header, bukan inferensi AI atau fakta terverifikasi.',
                         'Record-tabel teramati bukan observasi harga aktif. Tabel berdampingan dapat memakai baris fisik yang sama.',
                         'Tidak mengonversi mata uang, satuan, KKK, tanggal ambigu, atau menyatukan identitas.']
        if not tables:report['notes'].append('Tidak ada tabel yang cukup jelas; isi tidak dipaksa menjadi transaksi. Periksa raw atau ekspor rentang tabel.')
        if len(json.dumps(report,ensure_ascii=False).encode())>700000:
            for table in tables:
                table['preview']=[]
                for column in table['columns']:column['samples']=[]
            report['status']='partial';report['notes'].append('Preview nilai dibatasi oleh anggaran ukuran laporan.')
    except Exception:
        report.update(status='failed',classification='unknown',sheets=[],notes=['Breakdown aman tidak berhasil: format, macro, enkripsi, atau batas pemrosesan perlu diperiksa. Raw tetap utuh.'])
        report['summary']={key:0 for key in report['summary']}
    return report
