"""Schema-on-read intake: keep authorized original bytes before attempting extraction."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import os
import re
import subprocess
import sys
import unicodedata
from pathlib import Path
from urllib.parse import urlsplit

from .db import dumps, now, uid

CATEGORIES = {'general', 'price', 'production', 'supply_chain', 'weather', 'quality', 'traceability'}
MAX_TEXT = 600_000
MAX_FILE = int(os.environ.get('RUBBER_MAX_FILE_BYTES', 10 * 1024 * 1024))
MAX_OWNER_BYTES = int(os.environ.get('RUBBER_OWNER_RAW_BYTES', 256 * 1024 * 1024))
MAX_GLOBAL_BYTES = int(os.environ.get('RUBBER_GLOBAL_RAW_BYTES', 1024 * 1024 * 1024))
MAX_DOCUMENTS = int(os.environ.get('RUBBER_MAX_DOCUMENTS', 10000))


def clean_filename(filename):
    text = unicodedata.normalize('NFKC', Path((filename or 'raw.bin').replace('\\', '/')).name)
    return re.sub(r'[\x00-\x1f\x7f]', '', text).strip()[:180] or 'raw.bin'


def safe_source_uri(value):
    if not value: return ''
    if len(value) > 2000: raise ValueError('URL sumber terlalu panjang.')
    u = urlsplit(value)
    if u.scheme not in ('https', 'http') or not u.hostname or u.username or u.password or u.fragment:
        raise ValueError('URL sumber harus HTTP(S), tanpa kredensial atau fragmen. URL ini hanya metadata, tidak diambil otomatis.')
    # Never encourage credentials in a stored URL.
    if re.search(r'(?i)(api[_-]?key|token|secret|password|signature|credential)=', u.query):
        raise ValueError('Jangan sertakan kredensial dalam URL sumber.')
    return value


def extract_bytes(data: bytes, filename: str, media_type: str):
    suffix = Path(filename).suffix.lower()
    warnings = []
    status = 'indexed'
    parser = 'text-v1'
    text = ''
    try:
        if suffix == '.pdf' or data.startswith(b'%PDF-'):
            from pypdf import PdfReader
            parser = 'pdf-text-v1'
            pdf = PdfReader(io.BytesIO(data), strict=True)
            if pdf.is_encrypted:
                return '', 'raw_only', parser, ['PDF terenkripsi disimpan utuh; tidak dibuka atau ditebak kata sandinya.']
            parts = []
            for i, page in enumerate(pdf.pages):
                if i >= 50 or sum(map(len, parts)) >= MAX_TEXT:
                    warnings.append('Indeks PDF dibatasi 50 halaman / 600.000 karakter. Semua byte asli tetap disimpan.')
                    status = 'partial'
                    break
                parts.append(f'\n[Halaman {i + 1}]\n' + (page.extract_text() or ''))
            text = '\n'.join(parts)
            if not any(x.strip() for x in parts) or len(re.sub(r'\[Halaman \d+\]', '', text).strip()) < 10:
                status = 'raw_only'
                warnings.append('Tidak ada teks yang memadai. PDF hasil pindai memerlukan OCR; OCR belum dijalankan.')
        elif suffix in ('.xlsx', '.csv', '.tsv'):
            parser = 'tabular-archive-v1'
            lines = []
            if suffix == '.xlsx':
                from openpyxl import load_workbook
                from zipfile import ZipFile
                with ZipFile(io.BytesIO(data)) as z:
                    if len(z.infolist()) > 2000 or sum(i.file_size for i in z.infolist()) > 60 * 1024 * 1024:
                        raise ValueError('Rentang XLSX terlalu besar untuk indeks aman.')
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
                try:
                    count = 0
                    for sheet in wb:
                        if sheet.sheet_state != 'visible':
                            warnings.append('Lembar tersembunyi tidak diindeks.'); status = 'partial'; continue
                        for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                            count += 1
                            if count > 10000 or sum(map(len, lines)) > MAX_TEXT:
                                status = 'partial'; break
                            lines.append(f'{sheet.title}:{n} ' + ' | '.join(str(v)[:4096] if v is not None else '' for v in row[:64]))
                            if len(row) > 64: status = 'partial'
                        if count > 10000: break
                finally: wb.close()
                warnings.append('Formula tidak dijalankan; hanya nilai cache. Tidak dinormalisasi sebagai pembelian.')
            else:
                encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
                decoded = data.decode(encoding)
                try: dialect = csv.Sniffer().sniff(decoded[:12000], delimiters=',;\t|')
                except csv.Error: dialect = csv.excel_tab if suffix == '.tsv' else csv.excel
                for n, row in enumerate(csv.reader(io.StringIO(decoded), dialect), 1):
                    if n > 10001 or sum(map(len, lines)) > MAX_TEXT:
                        status = 'partial'; break
                    lines.append(f'Baris {n}: ' + ' | '.join(v[:4096] for v in row[:64]))
                    if len(row) > 64: status = 'partial'
            text = '\n'.join(lines)
            if status == 'partial': warnings.append('Indeks tabular dibatasi 10.000 baris, 64 kolom, dan 600.000 karakter; raw tidak dipotong.')
        elif suffix in ('.json', '.geojson') or media_type.split(';')[0] in ('application/json', 'application/geo+json'):
            parser = 'json-v1'
            def invalid_constant(_): raise ValueError('JSON berisi angka non-finite.')
            obj = json.loads(data, parse_constant=invalid_constant)
            text = json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False)
        elif suffix in ('.ndjson', '.jsonl') or media_type.split(';')[0] in ('application/x-ndjson',):
            parser = 'ndjson-v1'; parts = []
            for i, line in enumerate(data.decode('utf-8-sig').splitlines()):
                if not line.strip(): continue
                if i >= 10000:
                    status = 'partial'; warnings.append('Indeks NDJSON dibatasi 10.000 baris.'); break
                json.loads(line, parse_constant=lambda _: (_ for _ in ()).throw(ValueError('Angka non-finite.')))
                parts.append(line)
                if sum(map(len, parts)) > MAX_TEXT: status = 'partial'; break
            text = '\n'.join(parts)
        elif suffix in ('.xml', '.rss', '.atom') or media_type.split(';')[0] in ('application/xml', 'text/xml', 'application/rss+xml', 'application/atom+xml'):
            from defusedxml import ElementTree as ET
            parser = 'xml-safe-v1'
            root = ET.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            parts = []
            for i, el in enumerate(root.iter()):
                if i >= 20000: status = 'partial'; warnings.append('Indeks XML dibatasi 20.000 elemen.'); break
                if el.text and el.text.strip(): parts.append(el.text.strip())
            text = '\n'.join(parts)
        elif suffix in ('.html', '.htm') or media_type.split(';')[0] == 'text/html':
            from bs4 import BeautifulSoup
            parser = 'html-text-v1'
            soup = BeautifulSoup(data, 'html.parser')
            for el in soup(['script', 'style', 'iframe', 'object', 'embed', 'noscript']): el.decompose()
            text = soup.get_text('\n', strip=True)
            warnings.append('HTML diperlakukan sebagai teks; skrip, tautan eksternal, dan lampiran tidak dijalankan.')
        elif suffix in ('.txt', '.md', '.log', '.yaml', '.yml') or media_type.startswith('text/'):
            encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            text = data.decode(encoding)
            if '\x00' in text: raise ValueError('Payload tampak biner; tidak diindeks sebagai teks.')
        else:
            parser = 'opaque-v1'; status = 'raw_only'
            warnings.append('Format disimpan sebagai raw tanpa ekstraksi. Tidak ada OCR, eksekusi kode, pembukaan arsip, atau inferensi isi otomatis.')
    except Exception as exc:
        status = 'extraction_failed'
        parser = parser or 'opaque-v1'
        # Do not leak paths or payloads from library exceptions.
        warnings.append('Ekstraksi aman tidak berhasil; byte asli tetap tersedia. Periksa format, enkripsi, atau struktur file.')
        text = ''
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT]; status = 'partial'
        warnings.append('Indeks teks dibatasi 600.000 karakter; raw tetap lengkap.')
    return text, status, parser, warnings


def indexer_environment():
    # Windows needs SystemRoot for side-by-side/runtime DLL resolution. Keep only
    # these OS paths; DB/S3 credentials must not be inherited by the indexer.
    env={'PATH':os.environ.get('PATH',''),'PYTHONUTF8':'1','PYTHONDONTWRITEBYTECODE':'1',
         'OPENBLAS_NUM_THREADS':'1','OMP_NUM_THREADS':'1'}
    for key in ('SystemRoot','WINDIR','TEMP','TMP'):
        if os.environ.get(key):env[key]=os.environ[key]
    return env


def isolated_extract(path: Path, filename: str, media_type: str, *, include_breakdown=False):
    try:
        result=subprocess.run([sys.executable,'-m','backend.extract_worker',str(path.resolve()),filename,media_type]+(['--with-breakdown'] if include_breakdown else []),
                              cwd=Path(__file__).resolve().parents[1],capture_output=False,stdout=subprocess.PIPE,stderr=subprocess.DEVNULL,
                              timeout=18,env=indexer_environment())
        if result.returncode!=0 or len(result.stdout)>4*1024*1024: raise ValueError('Indexer budget exceeded.')
        obj=json.loads(result.stdout)
        values=(obj['text'],obj['status'],obj['parser'],obj['warnings'])
        return (*values,obj.get('breakdown')) if include_breakdown else values
    except (subprocess.TimeoutExpired,ValueError,OSError,KeyError):
        values=('', 'extraction_failed', 'isolated-indexer-v1', ['Batas waktu/memori/CPU indexer tercapai atau ekstraksi gagal. Raw tetap utuh; proses utama tidak menjalankan lampiran.'])
        return (*values,None) if include_breakdown else values


def isolated_breakdown(path: Path, filename: str):
    try:
        result=subprocess.run([sys.executable,'-m','backend.extract_worker',str(path.resolve()),filename,'','--breakdown-only'],
                              cwd=Path(__file__).resolve().parents[1],stdout=subprocess.PIPE,stderr=subprocess.DEVNULL,
                              timeout=18,env=indexer_environment())
        if result.returncode!=0 or len(result.stdout)>1024*1024:raise ValueError('Breakdown budget exceeded.')
        return json.loads(result.stdout)['breakdown']
    except (subprocess.TimeoutExpired,ValueError,OSError,KeyError):return None


class ArchiveMixin:
    def archive(self, owner, filename, data, *, category='general', context='', source_uri='', media_type='', logical_key=None,
                source_id=None, run_id=None, request_key=None, internal=False):
        from .service import public_version
        filename = clean_filename(filename)
        if category not in CATEGORIES: raise ValueError('Kategori karet tidak dikenali.')
        if len(context) > 2000: raise ValueError('Konteks maksimum 2.000 karakter.')
        if not data: raise ValueError('Payload kosong.')
        limit = 64 * 1024 * 1024 if internal else MAX_FILE
        if len(data) > limit: raise ValueError(f'Payload melampaui batas {limit // 1024 // 1024} MB. Tidak ada byte yang diterima sebagai arsip.')
        source_uri = safe_source_uri(source_uri)
        if request_key and not re.fullmatch(r'[A-Za-z0-9_.:-]{8,160}', request_key):
            raise ValueError('Idempotency-Key harus 8–160 karakter aman.')
        sha = hashlib.sha256(data).hexdigest()
        key = logical_key or 'archive:' + filename.casefold()
        received = now()
        meta = {'intake_kind': 'external' if source_id else 'archive', 'category': category, 'context': context,
                'source_uri': source_uri, 'media_type': (media_type or mimetypes.guess_type(filename)[0] or 'application/octet-stream')[:200],
                'received_at': received, 'scope': 'natural_rubber_declared', 'source_id': source_id, 'run_id': run_id,
                'rights': 'source_policy' if internal else 'user_confirmed', 'archive_contract': 'raw-v2'}
        with self.db.connect() as conn:
            conn.execute('BEGIN IMMEDIATE')
            if request_key:
                previous = conn.execute('SELECT r.*,d.logical_key FROM intake_requests r JOIN document_versions v ON v.id=r.version_id JOIN documents d ON d.id=v.document_id WHERE r.owner_id=? AND r.request_key=?', (owner, request_key)).fetchone()
                if previous:
                    if previous['sha256'] != sha or previous['logical_key'] != key:
                        from fastapi import HTTPException
                        raise HTTPException(409, 'Idempotency-Key sudah dipakai untuk byte berbeda.')
                    self.db.objects.put(f'{owner}/{sha}',data)
                    acquisition_id=uid()
                    conn.execute('INSERT INTO acquisitions VALUES(?,?,?,?,?,?)', (acquisition_id, owner, previous['version_id'], filename, received, 1))
                    conn.execute('INSERT INTO acquisition_details VALUES(?,?)',(acquisition_id,dumps(meta)))
                    return {'version': public_version(conn, previous['version_id']), 'duplicate': True, 'idempotent_replay': True}
            doc = conn.execute('SELECT * FROM documents WHERE owner_id=? AND logical_key=?', (owner, key)).fetchone()
            if doc:
                previous = conn.execute('SELECT v.id FROM document_versions v JOIN raw_assets a ON a.id=v.asset_id WHERE v.document_id=? AND a.sha256=? ORDER BY v.version_no DESC LIMIT 1', (doc['id'], sha)).fetchone()
            else: previous = None
            if previous:
                self.db.objects.put(f'{owner}/{sha}',data)
                vid = previous['id']
                acquisition_id=uid()
                conn.execute('INSERT INTO acquisitions VALUES(?,?,?,?,?,?)', (acquisition_id, owner, vid, filename, received, 1))
                conn.execute('INSERT INTO acquisition_details VALUES(?,?)',(acquisition_id,dumps(meta)))
                conn.execute('UPDATE documents SET latest_version_id=? WHERE id=?',(vid,doc['id']))
                conn.execute("UPDATE documents SET active_version_id=? WHERE id=? AND EXISTS(SELECT 1 FROM document_versions WHERE id=? AND status='archived')",(vid,doc['id'],vid))
                if request_key: conn.execute('INSERT INTO intake_requests VALUES(?,?,?,?,?)', (owner, request_key, sha, vid, received))
                self.db.activity(conn, owner, 'raw_duplicate', 'Akuisisi raw identik dicatat tanpa menggandakan byte.', doc['id'])
                return {'version': public_version(conn, vid), 'duplicate': True}
            asset = conn.execute('SELECT * FROM raw_assets WHERE owner_id=? AND sha256=?', (owner, sha)).fetchone()
            if not asset:
                size = conn.execute('SELECT COALESCE(SUM(size_bytes),0) FROM raw_assets WHERE owner_id=?', (owner,)).fetchone()[0]
                total = conn.execute('SELECT COALESCE(SUM(size_bytes),0) FROM raw_assets').fetchone()[0]
                if size + len(data) > MAX_OWNER_BYTES or total + len(data) > MAX_GLOBAL_BYTES:
                    raise ValueError('Kuota raw tercapai. Payload ditolak secara eksplisit, bukan dipotong. Atur kapasitas/retensi pada server.')
            if not doc:
                count = conn.execute('SELECT COUNT(*) FROM documents WHERE owner_id=?', (owner,)).fetchone()[0]
                if count >= MAX_DOCUMENTS: raise ValueError('Batas dokumen tercapai; atur kapasitas atau retensi pada server.')
                doc_id = uid()
                conn.execute('INSERT INTO documents(id,owner_id,filename,logical_key,created_at) VALUES(?,?,?,?,?)', (doc_id, owner, filename, key, received))
                number = 1
            else:
                doc_id = doc['id']
                number = conn.execute('SELECT COALESCE(MAX(version_no),0)+1 FROM document_versions WHERE document_id=?', (doc_id,)).fetchone()[0]
            if not asset:
                asset_id = uid(); object_key = f'{owner}/{sha}'
                self.db.objects.put(object_key, data)
                conn.execute('INSERT INTO raw_assets(id,owner_id,sha256,object_key,size_bytes) VALUES(?,?,?,?,?)', (asset_id, owner, sha, object_key, len(data)))
            else:
                self.db.objects.put(asset['object_key'],data)
                asset_id = asset['id']
            vid = uid()
            conn.execute('INSERT INTO document_versions(id,document_id,asset_id,version_no,status,created_at,intake_meta) VALUES(?,?,?,?,?,?,?)',
                         (vid, doc_id, asset_id, number, 'queued', received, dumps(meta)))
            conn.execute('UPDATE documents SET latest_version_id=? WHERE id=?', (vid, doc_id))
            acquisition_id=uid()
            conn.execute('INSERT INTO acquisitions VALUES(?,?,?,?,?,?)', (acquisition_id, owner, vid, filename, received, 0))
            conn.execute('INSERT INTO acquisition_details VALUES(?,?)',(acquisition_id,dumps(meta)))
            conn.execute('INSERT INTO jobs(id,version_id,kind,status,created_at) VALUES(?,?,?,?,?)', (uid(), vid, 'extract', 'queued', received))
            if request_key: conn.execute('INSERT INTO intake_requests VALUES(?,?,?,?,?)', (owner, request_key, sha, vid, received))
            self.db.activity(conn, owner, 'raw_received', f'Raw karet diterima: {filename}', doc_id, {'sha256': sha, 'category': category, 'version_no': number})
            return {'version': public_version(conn, vid), 'duplicate': False}

    def _extract(self, job):
        with self.db.connect() as conn:
            row = conn.execute('SELECT a.object_key,a.sha256,d.id AS document_id,d.filename,d.owner_id,v.intake_meta FROM document_versions v JOIN raw_assets a ON a.id=v.asset_id JOIN documents d ON d.id=v.document_id WHERE v.id=?', (job['version_id'],)).fetchone()
            if not row: return
        meta = json.loads(row['intake_meta'])
        path = self.db.object_path(row['object_key'], row['sha256'])
        text, status, parser, warnings, breakdown = isolated_extract(path, row['filename'], meta.get('media_type', ''), include_breakdown=True)
        parsed = {'sheets': [], 'warnings': warnings, 'parser_version': parser, 'extraction_status': status, 'indexed_characters': len(text), 'row_count': 0}
        if breakdown is None:
            from .breakdown_service import failed_breakdown
            breakdown=failed_breakdown()
        parsed['breakdown']=breakdown
        with self.db.connect() as conn:
            conn.execute('DELETE FROM archive_content WHERE version_id=?', (job['version_id'],))
            conn.execute('INSERT INTO archive_content VALUES(?,?,?,?,?,?)', (job['version_id'], status, parser, text, dumps(warnings), now()))
            conn.execute('UPDATE document_versions SET status=?,parsed_meta=?,error=NULL WHERE id=?', ('archived', dumps(parsed), job['version_id']))
            conn.execute('UPDATE documents SET active_version_id=? WHERE id=? AND latest_version_id=?', (job['version_id'], row['document_id'], job['version_id']))
            conn.execute('UPDATE jobs SET status=?,finished_at=? WHERE id=?', ('done', now(), job['id']))
            self.db.activity(conn, row['owner_id'], 'raw_indexed', 'Arsip raw tersimpan; status ekstraksi: ' + status + '.', row['document_id'], {'warnings': warnings})

    def archive_search(self, owner, q='', limit=20, offset=0):
        escaped = q.lower().replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
        where = " FROM documents d JOIN document_versions v ON v.id=d.active_version_id JOIN archive_content c ON c.version_id=v.id WHERE d.owner_id=? "
        params = [owner]
        if q:
            if self.db.kind=='postgresql':
                where += " AND (LOWER(d.filename) LIKE ? ESCAPE '\\' OR to_tsvector('simple',c.text_content) @@ plainto_tsquery('simple',?) OR LOWER(v.intake_meta) LIKE ? ESCAPE '\\')"
                params += [f'%{escaped}%',q,f'%{escaped}%']
            else:
                where += " AND (LOWER(d.filename) LIKE ? ESCAPE '\\' OR LOWER(c.text_content) LIKE ? ESCAPE '\\' OR LOWER(v.intake_meta) LIKE ? ESCAPE '\\')"
                params += [f'%{escaped}%'] * 3
        with self.db.connect() as conn:
            total = conn.execute('SELECT COUNT(*)' + where, params).fetchone()[0]
            rows = conn.execute('SELECT d.id AS document_id,d.filename,v.id AS version_id,v.intake_meta,c.extraction_status,c.parser,c.text_content,c.warnings' + where + ' ORDER BY d.created_at DESC LIMIT ? OFFSET ?', params + [limit, offset]).fetchall()
            items = []
            for r in rows:
                d = dict(r); text = d.pop('text_content')
                at = max(0, text.lower().find(q.lower()) - 140) if q else 0
                d['excerpt'] = text[at:at + 1200]
                d['intake_meta'] = json.loads(d['intake_meta']); d['warnings'] = json.loads(d['warnings'])
                items.append(d)
            return {'items': items, 'total': total, 'offset': offset, 'limit': limit}

    def acquisitions(self, version_id, owner, limit=100, offset=0):
        with self.db.connect() as conn:
            self.can_read(conn, version_id, owner)
            total = conn.execute('SELECT COUNT(*) FROM acquisitions WHERE version_id=?', (version_id,)).fetchone()[0]
            rows = conn.execute('SELECT a.id,a.filename,a.captured_at,a.duplicate,d.metadata FROM acquisitions a LEFT JOIN acquisition_details d ON d.acquisition_id=a.id WHERE a.version_id=? ORDER BY a.captured_at DESC LIMIT ? OFFSET ?', (version_id, limit, offset)).fetchall()
            return {'items': [{**dict(r),'metadata':json.loads(r['metadata']) if r['metadata'] else {}} for r in rows], 'total': total, 'limit': limit, 'offset': offset}

    def archive_content(self, version_id, owner, offset=0, limit=12000):
        with self.db.connect() as conn:
            self.can_read(conn, version_id, owner)
            row=conn.execute('SELECT extraction_status,parser,warnings,SUBSTR(text_content,?,?) AS text,LENGTH(text_content) AS total_characters FROM archive_content WHERE version_id=?',(offset+1,limit,version_id)).fetchone()
            if not row: raise LookupError('Indeks raw belum tersedia; file asli tetap dapat diunduh.')
            return {**dict(row),'warnings':json.loads(row['warnings']),'offset':offset,'limit':limit}
